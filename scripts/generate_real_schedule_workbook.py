# This file is mostly AI generated.
"""Build the plain November 2025 ward intake workbook from companion data."""

import argparse
import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from ruamel.yaml import YAML

SOURCE = (
    Path(__file__).resolve().parents[1]
    / "docs/content/user-guide/build-a-real-schedule"
)
DEFAULT_OUTPUT = SOURCE / "unfilled-ward-schedule-2025-11.xlsx"
FREEDAY_FILL = PatternFill("solid", fgColor="FFEAF0F7")
SENIOR_FILL = PatternFill("solid", fgColor="FFFFF0BC")
STRONG_FONT = Font(name="Calibri", size=11, color="FFC6252F")
NORMAL_FONT = Font(name="Calibri", size=11, color="FF202A35")


def request_rows(filename):
    with (SOURCE / filename).open(newline="", encoding="utf-8-sig") as stream:
        return {row[0]: row[1:] for row in csv.reader(stream)}


people = (SOURCE / "people.txt").read_text(encoding="utf-8").splitlines()
with (SOURCE / "people-history.csv").open(newline="", encoding="utf-8-sig") as stream:
    history = {
        person: (shift, int(count)) for person, shift, count in csv.reader(stream)
    }
strong = request_rows("shift-requests-strong.csv")
moderate = request_rows("shift-requests-moderate.csv")
reference = YAML(typ="safe").load(
    (SOURCE / "reference.yaml").read_text(encoding="utf-8")
)
groups = {item["id"]: set(item["members"]) for item in reference["peopleGroups"]}
freedays = {
    int(day)
    for item in reference["dateGroups"]
    if item["id"] == "FREEDAY"
    for day in item["members"]
}
assert len(people) == len(history) == len(strong) == len(moderate) == 87
assert all(len(strong[person]) == len(moderate[person]) == 30 for person in people)


def compact_history(person):
    shift, count = history[person]
    code = "off" if shift == "OFF" else shift
    return f"{count}{code}" if count > 1 else code


def role(person):
    if person == "P1":
        return "HN"
    if person == "P2":
        return "AHN"
    return "N/A" if person in groups["Students"] else "N"


book = Workbook()
sheet = book.active
sheet.title = "November 2025"
sheet["A1"] = "November 2025"
sheet["A1"].font = Font(name="Calibri", size=12, bold=True)
for column, label in enumerate(("Role", "History", "Name"), 1):
    sheet.cell(2, column, label)
for day in range(1, 31):
    column = day + 3
    sheet.cell(2, column, day)
    sheet.cell(3, column, date(2025, 11, day).strftime("%a"))
    sheet.column_dimensions[sheet.cell(2, column).column_letter].width = 6
    if day in freedays:
        sheet.cell(2, column).fill = FREEDAY_FILL
        sheet.cell(3, column).fill = FREEDAY_FILL

sheet.column_dimensions["A"].width = 9
sheet.column_dimensions["B"].width = 10
sheet.column_dimensions["C"].width = 10
sheet.freeze_panes = "D4"
sheet.print_title_rows = "1:3"

ordered = ["P1", "P2"]
for group in ("Day People", "Evening People", "Night People"):
    members = [
        person for person in people if person in groups[group] and person not in ordered
    ]
    ordered.extend(members)
    if group != "Night People":
        ordered.extend((None, None))
assert [person for person in ordered if person] == people

strong_count = moderate_count = 0
for row, person in enumerate(ordered, 4):
    if person is None:
        continue
    sheet.cell(row, 1, role(person))
    sheet.cell(row, 2, compact_history(person))
    name = sheet.cell(row, 3, person)
    if person in groups["Senior Nurses"]:
        name.fill = SENIOR_FILL
    for day in range(1, 31):
        strong_value = strong[person][day - 1]
        moderate_value = moderate[person][day - 1]
        assert not (strong_value and moderate_value)
        source_value = strong_value or moderate_value
        cell = sheet.cell(row, day + 3)
        if day in freedays:
            cell.fill = FREEDAY_FILL
        if not source_value:
            continue
        cell.value = "1" if source_value == "OFF" else source_value
        cell.font = STRONG_FONT if strong_value else NORMAL_FONT
        strong_count += bool(strong_value)
        moderate_count += bool(moderate_value)

assert strong_count == 408 and moderate_count == 90
book.properties.title = "November 2025"
book.properties.subject = "Unfilled anonymized ward schedule"
book.properties.creator = "Nurse Scheduling Project"
book.properties.keywords = "This file is mostly AI generated."
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
output = parser.parse_args().output
book.save(output)
book.close()
print(output)
