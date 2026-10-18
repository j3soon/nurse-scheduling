"""Inspect optimizer assignments, signed requests and roster statistics without parsing annotations as shifts."""

# This file is part of Nurse Scheduling Project, see <https://github.com/j3soon/nurse-scheduling>.
#
# Copyright (C) 2023-2026 Johnson Sun
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as
# published by the Free Software Foundation, either version 3 of the
# License, or (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.

# This file is mostly AI generated.

import argparse
import hashlib
import json
import re
import sys
from collections import Counter
from collections.abc import Sequence
from decimal import Decimal, InvalidOperation
from itertools import islice, takewhile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from ruamel.yaml import YAML

try:
    from .inspect_xlsx import _check_archive
except ImportError:  # Standalone copy in /reference/tools.
    from inspect_xlsx import _check_archive

RESULT = "/workspace/optimizer-results/optimized-schedule.xlsx"
CONTEXT = "/workspace/optimizer-results/schedule-context.json"
MAX_CELLS = 1_000_000


def _weight(value: Any) -> Decimal:
    if value == ".inf":
        return Decimal("Infinity")
    if value == "-.inf":
        return Decimal("-Infinity")
    try:
        result = Decimal(str(value))
    except InvalidOperation as error:
        raise ValueError(f"Invalid numeric request weight: {value}") from error
    if result.is_nan():
        raise ValueError("NaN is not a request weight")
    return result


def _assigned(value: Any, shift_types: list[str]) -> str:
    text = "" if value is None else str(value)
    annotations = r"(?:\s*\[[^\]]*\])*\s*"
    candidates = []
    if re.fullmatch(annotations, text):
        candidates.append("OFF")
    for shift in shift_types:
        if text.startswith(shift) and re.fullmatch(annotations, text[len(shift) :]):
            candidates.append(shift)
    if len(candidates) != 1:
        raise ValueError(f"Assignment is unknown or ambiguous: {value!r}")
    return candidates[0]


def _date_columns(row, dates: list[str]) -> dict[str, int] | None:
    columns = {}
    for date in dates:
        year, month, day = map(int, date.split("-"))
        # The exporter uses day, month/day, or year/month/day according to range.
        labels = {date, str(day), f"{month}/{day}", f"{year}/{month}/{day}"}
        matches = [col for col, cell in enumerate(row) if str(cell.value) in labels]
        if len(matches) != 1:
            return None
        columns[date] = matches[0]
    return columns if len(set(columns.values())) == len(dates) else None


def _read_assignments(path: Path, context: dict) -> tuple[dict, Any, Any]:
    _check_archive(path)
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        candidates = []
        for sheet in workbook:
            if sheet.max_row * sheet.max_column > MAX_CELLS:
                raise ValueError("Optimizer worksheet is too large")
            rows = list(sheet.iter_rows())
            people = set(context["people"])
            for header, row in enumerate(rows[:10]):
                columns = _date_columns(row, context["dates"])
                if columns is None:
                    continue
                assigned = {}
                last_person_row = header
                for row_index, cells in enumerate(rows[header + 1 :], header + 1):
                    person = cells[0].value
                    if person not in people:
                        continue
                    if person in assigned:
                        raise ValueError(f"Duplicate person row: {person}")
                    assigned[person] = {}
                    for date, col in columns.items():
                        cell = cells[col]
                        if cell.data_type == "f":
                            raise ValueError("Formula assignments are not supported")
                        assigned[person][date] = _assigned(cell.value, context["shift_types"])
                    last_person_row = row_index
                    if set(assigned) == people:
                        break
                if set(assigned) != people:
                    continue
                metadata = {}
                for cells in rows[last_person_row + 1 :]:
                    if cells[0].value in {"Score", "Status"}:
                        values = [c.value for c in cells[1:] if c.value is not None]
                        if len(values) == 1:
                            metadata[cells[0].value] = values[0]
                candidates.append((assigned, metadata.get("Score"), metadata.get("Status")))
        if len(candidates) != 1:
            raise ValueError("Could not uniquely locate the optimizer roster and date headers")
        return candidates[0]
    finally:
        workbook.close()


def inspect_result(
    path: Path,
    context: dict,
    source_sha256: str,
    weights: list[str] | None = None,
    max_unmet: int = 20,
    *,
    people: list[str] | None = None,
    dates: list[str] | None = None,
    max_assignments: int = 100,
) -> dict:
    """Audit nonzero-weight person/date requests, independently of export markers."""
    if context.get("schema_version") != 1 or context.get("source_sha256") != source_sha256:
        raise ValueError("The compiled context does not match the optimizer source")
    if max_unmet < 0:
        raise ValueError("max_unmet must not be negative")
    assignments, score, status = _read_assignments(path, context)
    if people is not None or dates is not None:
        if not 1 <= max_assignments <= 1000:
            raise ValueError("max_assignments must be between 1 and 1000")
        person_ids = {str(person): person for person in context["people"]}
        if len(person_ids) != len(context["people"]):
            raise ValueError("Person IDs cannot be uniquely selected as text")
        selected_people = list(dict.fromkeys(people)) if people else list(person_ids)
        selected_dates = list(dict.fromkeys(dates)) if dates else context["dates"]
        if set(selected_people) - person_ids.keys():
            raise ValueError("Unknown person ID in assignment query")
        if set(selected_dates) - set(context["dates"]):
            raise ValueError("Unknown ISO date in assignment query")
        rows = (
            {"person": person_ids[person], "date": date, "shift_type": assignments[person_ids[person]][date]}
            for person in selected_people
            for date in selected_dates
        )
        return {
            "source_sha256": source_sha256,
            "assignments": list(islice(rows, max_assignments)),
            "assignments_truncated": len(selected_people) * len(selected_dates) > max_assignments,
        }
    selected = {_weight(w) for w in weights} if weights else None
    summary = {}
    unmet = []
    for request in context["requests"]:
        weight = _weight(request["weight"])
        if selected is not None and weight not in selected:
            continue
        row = summary.setdefault(weight, {"weight": request["weight"], "total": 0, "satisfied": 0, "unmet": 0})
        for person in request["people"]:
            for date in request["dates"]:
                assigned = assignments[person][date]
                matches = assigned in request["shift_types"]
                satisfied = matches if weight > 0 else not matches
                row["total"] += 1
                row["satisfied" if satisfied else "unmet"] += 1
                if not satisfied:
                    unmet.append(
                        {
                            "preference_index": request["preference_index"],
                            "person": person,
                            "date": date,
                            "requested": request["shift_types"],
                            "assigned": assigned,
                            "weight": request["weight"],
                        }
                    )
    result = {
        "source_sha256": source_sha256,
        "score": score,
        "score_direction": "maximize",
        "status": status,
        "summary": list(summary.values()),
        "unmet": unmet[:max_unmet],
        "unmet_truncated": len(unmet) > max_unmet,
    }

    if "policy" in context:
        if context.get("workbook_sha256") == hashlib.sha256(path.read_bytes()).hexdigest():
            result["policy"] = context["policy"]
            result["policy_scope"] = context["policy_scope"]
        else:
            result["policy_unavailable"] = "The workbook does not match the policy audit context."
    return result


def _longest_run(shifts: Sequence[str], working: bool, history: Sequence[str] = ()) -> int:
    matches = lambda shift: (shift != "OFF") == working
    run = sum(1 for _ in takewhile(matches, reversed(history)))
    longest = 0
    for shift in shifts:
        run = run + 1 if matches(shift) else 0
        longest = max(longest, run)
    return longest


def _statistics(assignments, context, people, histories, include_history, max_people, summary_only):
    person_ids = {str(person): person for person in context["people"]}
    if len(person_ids) != len(context["people"]):
        raise ValueError("Person IDs cannot be uniquely selected as text")
    selected = list(dict.fromkeys(people)) if people else list(person_ids)
    if set(selected) - person_ids.keys():
        raise ValueError("Unknown person ID in statistics query")
    rows = []
    for selected_id in selected:
        person = person_ids[selected_id]
        shifts = [assignments[person][day] for day in context["dates"]]
        row = {
            "person": person,
            "shift_counts": dict(Counter(shifts)),
            "max_work_run": _longest_run(shifts, True),
            "max_off_run": _longest_run(shifts, False),
        }
        if include_history:
            history = histories.get(person)
            row.update(
                history_days=None if history is None else len(history),
                max_work_run_including_history=None if history is None else _longest_run(shifts, True, history),
                max_off_run_including_history=None if history is None else _longest_run(shifts, False, history),
            )
        rows.append(row)
    shift_counts = Counter()
    for row in rows:
        shift_counts.update(row["shift_counts"])
    summary = {"people": len(rows), "shift_counts": dict(shift_counts)}
    run_fields = ["max_work_run", "max_off_run"]
    if include_history:
        run_fields += ["max_work_run_including_history", "max_off_run_including_history"]
    for field in run_fields:
        known = [row[field] for row in rows if row[field] is not None]
        summary[field] = max(known, default=None)
        summary[field + "_distribution"] = dict(sorted(Counter(known).items()))
        if include_history and field.endswith("including_history"):
            summary[field + "_unknown_people"] = len(rows) - len(known)
    return {
        "source_sha256": context["source_sha256"],
        "scope": "Observed assignments over the full scheduled date range, not a policy audit. "
        "Counts exclude history. Extended runs include only history suffixes crossing the first scheduled date. "
        "Unknown history remains null.",
        "dates": context["dates"],
        "summary": summary,
        "statistics": [] if summary_only else rows[:max_people],
        "statistics_truncated": bool(rows) if summary_only else len(rows) > max_people,
    }


def summarize_roster(
    path, context, source_sha256, *, people=None, max_people=100, include_history=False, source=None, summary_only=False
):
    """Summarize observed assignments, independently of scheduling policy."""
    if context.get("schema_version") != 1 or context.get("source_sha256") != source_sha256:
        raise ValueError("The compiled context does not match the optimizer source")
    if not 1 <= max_people <= 1000:
        raise ValueError("max_people must be between 1 and 1000")
    histories = {}
    if include_history and source is not None and source.is_file():
        content = source.read_bytes()
        if hashlib.sha256(content).hexdigest() == source_sha256:
            data = YAML(typ="safe").load(content)
            histories = {person["id"]: person.get("history") for person in data["people"]["items"]}
    assignments, _, _ = _read_assignments(path, context)
    return _statistics(assignments, context, people, histories, include_history, max_people, summary_only)


def _weight_arguments(arguments: list[str]) -> list[str]:
    """Keep negative numeric weights from being mistaken for option names."""
    normalized = []
    index = 0
    while index < len(arguments):
        if arguments[index] == "--weight" and index + 1 < len(arguments):
            value = arguments[index + 1]
            try:
                _weight(value)
            except ValueError:
                pass
            else:
                normalized.append(f"--weight={value}")
                index += 2
                continue
        normalized.append(arguments[index])
        index += 1
    return normalized


def load_context(source_sha256: str, path: Path | None = None) -> dict:
    """Resolve the approved or pending context by source, respecting an explicit path."""
    if path is not None:
        return json.loads(path.read_text(encoding="utf-8"))
    approved = Path(CONTEXT)
    for candidate in (approved, approved.with_name("pending-schedule-context.json")):
        try:
            context = json.loads(candidate.read_text(encoding="utf-8"))
        except (FileNotFoundError, ValueError):
            continue
        if (
            isinstance(context, dict)
            and context.get("schema_version") == 1
            and context.get("source_sha256") == source_sha256
        ):
            return context
    raise ValueError("No hydrated optimizer context matches the completion source")


def main() -> None:
    parser = argparse.ArgumentParser(description="Inspect optimizer assignments, request audits and roster statistics.")
    parser.add_argument("workbook", nargs="?", type=Path, default=Path(RESULT))
    parser.add_argument("--context", type=Path)
    parser.add_argument("--source-sha256", required=True, help="source_sha256 from the optimizer completion")
    parser.add_argument(
        "--statistics", action="store_true", help="Count shifts and longest work/OFF runs, not policy compliance"
    )
    parser.add_argument(
        "--weight", action="append", help="numeric request weight, repeatable. Use --weight=-.inf for bans"
    )
    parser.add_argument("--max-unmet", type=int)
    parser.add_argument(
        "--person", action="append", help="Exact person ID, repeatable. Selects assignments unless --statistics is set."
    )
    parser.add_argument(
        "--date", action="append", help="ISO date YYYY-MM-DD, repeatable. Selects assignment-query output."
    )
    parser.add_argument("--max-assignments", type=int, help="Assignment-query row limit, at most 1000. Default: 100.")
    parser.add_argument("--summary-only", action="store_true", help="Return statistics totals without person rows")
    parser.add_argument("--max-people", type=int, help="Statistics row limit, at most 1000. Default: 100.")
    parser.add_argument("--include-history", action="store_true", help="Include known history suffixes in statistics")
    parser.add_argument(
        "--source",
        type=Path,
        help="Statistics history source. Defaults to /workspace/schedule.yaml. Used only when its bytes match source_sha256.",
    )
    args = parser.parse_args(_weight_arguments(sys.argv[1:]))
    if args.statistics:
        if any(value is not None for value in (args.weight, args.max_unmet, args.date, args.max_assignments)):
            parser.error("--statistics cannot be combined with --weight, --max-unmet, --date or --max-assignments")
    elif args.summary_only or args.max_people is not None or args.include_history or args.source is not None:
        parser.error("--summary-only, --max-people, --include-history and --source require --statistics")
    try:
        context = load_context(args.source_sha256, args.context)
        if args.statistics:
            result = summarize_roster(
                args.workbook,
                context,
                args.source_sha256,
                people=args.person,
                max_people=100 if args.max_people is None else args.max_people,
                include_history=args.include_history,
                source=args.source or Path("/workspace/schedule.yaml"),
                summary_only=args.summary_only,
            )
        else:
            result = inspect_result(
                args.workbook,
                context,
                args.source_sha256,
                args.weight,
                20 if args.max_unmet is None else args.max_unmet,
                people=args.person,
                dates=args.date,
                max_assignments=100 if args.max_assignments is None else args.max_assignments,
            )
    except (ValueError, KeyError, OSError) as error:
        parser.exit(1, f"{error}\n")
    print(json.dumps(result, ensure_ascii=False, allow_nan=False))


if __name__ == "__main__":
    main()
