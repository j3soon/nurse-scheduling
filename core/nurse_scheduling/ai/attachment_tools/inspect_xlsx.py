"""Print bounded, sheet-aware content from an XLSX workbook."""

# This file is part of Nurse Scheduling Project, see <https://github.com/j3soon/nurse-scheduling>.
#
# Copyright (C) 2023-2026 Johnson Sun
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as
# published by the Free Software Foundation, either version 3 of the
# License, or (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.

# This code is mostly AI generated.

import argparse
import json
from collections import defaultdict
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from defusedxml import ElementTree
from openpyxl import load_workbook
from openpyxl.cell.read_only import ReadOnlyCell
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.xml.functions import DEFUSEDXML

MAX_ARCHIVE_ENTRIES = 1_000
MAX_UNCOMPRESSED_BYTES = 50_000_000
MAX_SHEETS = 20

if not DEFUSEDXML:
    raise RuntimeError("defusedxml is required for untrusted XLSX files")


def _check_archive(path: Path) -> None:
    """Reject oversized or malformed workbook archives before XML parsing."""
    with path.open("rb") as stream:
        signature = stream.read(4)
    if signature != b"PK\x03\x04":
        raise ValueError("Not an XLSX archive")
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise ValueError("XLSX archive has too many entries")
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise ValueError("Encrypted XLSX files are not supported")
            if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("XLSX archive expands beyond the allowed size")
            names = {entry.filename for entry in entries}
            if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
                raise ValueError("Not an XLSX workbook")
    except BadZipFile as exc:
        raise ValueError("Invalid XLSX archive") from exc


def _theme_colors(workbook: Any) -> list[str | None]:
    if not workbook.loaded_theme:
        return []
    root = ElementTree.fromstring(workbook.loaded_theme)
    namespace = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    scheme = root.find(f"{namespace}themeElements/{namespace}clrScheme")
    if scheme is None:
        return []
    colors = []
    for name in (
        "lt1",
        "dk1",
        "lt2",
        "dk2",
        "accent1",
        "accent2",
        "accent3",
        "accent4",
        "accent5",
        "accent6",
        "hlink",
        "folHlink",
    ):
        entry = scheme.find(f"{namespace}{name}")
        node = next(iter(entry)) if entry is not None and len(entry) else None
        colors.append((node.get("lastClr") or node.get("val")) if node is not None else None)
    return colors


def _stored_color(color: Any, workbook: Any, theme: list[str | None]) -> dict[str, Any] | None:
    if color is None:
        return None
    result = {"type": color.type, "value": color.value}
    if color.tint:
        result["tint"] = color.tint
    if color.type == "theme" and 0 <= color.theme < len(theme):
        result["base_rgb"] = theme[color.theme]
    elif color.type == "indexed" and 0 <= color.indexed < len(workbook._colors):
        result["base_rgb"] = workbook._colors[color.indexed]
    return result


def _stored_style(cell: Any, workbook: Any, theme: list[str | None]) -> dict[str, Any]:
    """Describe stored styles, without evaluating conditional formatting or tint."""
    font = cell.font
    fill = cell.fill
    color = lambda value: _stored_color(value, workbook, theme)
    result = {
        "font": {
            key: value
            for key, value in {
                "name": font.name,
                "size": font.sz,
                "bold": font.b,
                "italic": font.i,
                "underline": font.u,
                "strike": font.strike,
                "color": color(font.color),
            }.items()
            if value not in (None, False)
        },
        "number_format": cell.number_format,
    }
    if getattr(fill, "patternType", None):
        result["fill"] = {
            "pattern": fill.patternType,
            "foreground": color(fill.fgColor),
            "background": color(fill.bgColor),
        }
    elif getattr(fill, "type", None):
        result["fill"] = {
            "gradient": fill.type,
            "stops": [{"position": stop.position, "color": color(stop.color)} for stop in fill.stop],
        }
    alignment = {
        key: getattr(cell.alignment, key)
        for key in (
            "horizontal",
            "vertical",
            "textRotation",
            "wrapText",
            "shrinkToFit",
            "indent",
            "readingOrder",
        )
        if getattr(cell.alignment, key) not in (None, False, 0)
    }
    if alignment:
        result["alignment"] = alignment
    borders = {}
    for key in ("left", "right", "top", "bottom", "diagonal", "start", "end", "vertical", "horizontal"):
        side = getattr(cell.border, key)
        if side is not None and (side.style or side.color):
            borders[key] = {"style": side.style, "color": color(side.color)}
    if borders:
        result["borders"] = borders
    return result


def _sheet_metadata(sheet: Any) -> dict[str, Any]:
    return {
        "name": sheet.title,
        "state": sheet.sheet_state,
        "reported_rows": sheet.max_row,
        "reported_columns": sheet.max_column,
    }


def inspect_workbook(
    path: Path,
    *,
    sheet_name: str | None = None,
    start_row: int = 1,
    start_column: int = 1,
    max_rows: int = 200,
    max_columns: int = 50,
    styles: bool = False,
    overview: bool = False,
) -> dict[str, Any]:
    """Return bounded cells with both formulas and last-saved cached values."""
    if min(start_row, start_column, max_rows, max_columns) <= 0:
        raise ValueError("Row and column positions and limits must be positive")
    if max_rows * max_columns > 10_000:
        raise ValueError("Select at most 10,000 cells per sheet")
    _check_archive(path)
    formula_book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    cached_book = None
    try:
        available = formula_book.sheetnames
        if sheet_name is not None and sheet_name not in available:
            raise ValueError(f"Unknown sheet {sheet_name!r}. Available sheets: {available}")
        selected = [sheet_name] if sheet_name is not None else available[: 100 if overview else MAX_SHEETS]
        if overview:
            return {
                "path": str(path),
                "sheet_count": len(available),
                "sheets": [_sheet_metadata(formula_book[name]) for name in selected],
                "sheets_truncated": sheet_name is None and len(available) > 100,
            }
        cached_book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        sheets = []
        style_table: dict[str, Any] = {}
        theme = _theme_colors(formula_book) if styles else []
        if styles and selected:
            style_table["0"] = _stored_style(ReadOnlyCell(formula_book[selected[0]], 1, 1, None), formula_book, theme)
        for name in selected:
            worksheet = formula_book[name]
            cached_sheet = cached_book[name]
            rows = []
            formula_rows = worksheet.iter_rows(
                min_row=start_row,
                max_row=start_row + max_rows - 1,
                min_col=start_column,
                max_col=start_column + max_columns - 1,
            )
            cached_rows = cached_sheet.iter_rows(
                min_row=start_row,
                max_row=start_row + max_rows - 1,
                min_col=start_column,
                max_col=start_column + max_columns - 1,
            )
            for row_number, (formula_row, cached_row) in enumerate(
                zip(formula_rows, cached_rows, strict=True), start=start_row
            ):
                values = []
                for formula_cell, cached_cell in zip(formula_row, cached_row, strict=True):
                    value = formula_cell.value
                    if formula_cell.data_type == "f":
                        value = {
                            "cell": formula_cell.coordinate,
                            "formula": getattr(value, "text", value),
                            "cached_value": cached_cell.value,
                        }
                    style_id = getattr(formula_cell, "_style_id", 0)
                    if styles and style_id:
                        key = str(style_id)
                        if key not in style_table:
                            style_table[key] = _stored_style(formula_cell, formula_book, theme)
                        if not isinstance(value, dict):
                            value = {"cell": formula_cell.coordinate, "value": value}
                        value["style"] = key
                    values.append(value)
                while values and values[-1] is None:
                    values.pop()
                if values:
                    rows.append({"row": row_number, "values": values})
            sheets.append(
                {
                    **_sheet_metadata(worksheet),
                    "first_column": start_column,
                    "rows": rows,
                    "truncated": (
                        start_row > 1
                        or start_column > 1
                        or (worksheet.max_row or 0) > start_row + max_rows - 1
                        or (worksheet.max_column or 0) > start_column + max_columns - 1
                    ),
                }
            )
        result = {
            "path": str(path),
            "sheet_names": available,
            "sheets_truncated": sheet_name is None and len(available) > MAX_SHEETS,
            "cached_values_are_last_saved_not_recalculated": True,
            "sheets": sheets,
        }
        if styles:
            result["styles"] = style_table
            result["cells_without_style_use"] = "0"
            result["styles_are_stored_not_rendered"] = True
        return result
    finally:
        formula_book.close()
        if cached_book is not None:
            cached_book.close()


def _count_cells(cells: list[tuple[str, Any]], styles: bool, max_groups: int) -> dict[str, Any]:
    groups = {}
    blanks = formulas = 0
    for coordinate, cell in cells:
        style = cell.get("style", "0") if isinstance(cell, dict) else "0"
        if isinstance(cell, dict) and "formula" in cell:
            descriptor = {"type": "formula", "formula": cell["formula"], "cached_value": cell["cached_value"]}
            formulas += 1
        else:
            value = cell.get("value") if isinstance(cell, dict) else cell
            if value is None:
                blanks += 1
            descriptor = {"type": type(value).__name__, "value": value}
        if styles:
            descriptor["style"] = style
        key = json.dumps(descriptor, sort_keys=True, default=str)
        group = groups.setdefault(key, {**descriptor, "count": 0, "first_cell": coordinate})
        group["count"] += 1
    ordered = sorted(groups.values(), key=lambda group: -group["count"])
    return {
        "cells": len(cells),
        "blank_cells": blanks,
        "formula_cells": formulas,
        "group_count": len(groups),
        "groups": ordered[:max_groups],
        "groups_truncated": len(groups) > max_groups,
    }


def summarize_workbook(
    path: Path,
    *,
    cell_range: str,
    sheet_name: str | None = None,
    axis: str = "all",
    styles: bool = False,
    max_groups: int = 100,
) -> dict[str, Any]:
    """Count exact typed values without recalculating or interpreting formulas."""
    if axis not in {"all", "row", "column"}:
        raise ValueError("axis must be all, row or column")
    if not 1 <= max_groups <= 1_000:
        raise ValueError("max_groups must be between 1 and 1,000")
    range_sheet, first_row, first_column, height, width = _parse_cell_range(cell_range)
    if range_sheet is not None:
        if sheet_name is not None and sheet_name != range_sheet:
            raise ValueError("--sheet conflicts with the sheet name in --range")
        sheet_name = range_sheet
    result = inspect_workbook(
        path,
        sheet_name=sheet_name,
        start_row=first_row,
        start_column=first_column,
        max_rows=height,
        max_columns=width,
        styles=styles,
    )
    for sheet in result["sheets"]:
        rows = {row["row"]: row["values"] for row in sheet.pop("rows")}
        cells = []
        buckets = defaultdict(list)
        for row_number in range(first_row, first_row + height):
            values = rows.get(row_number, [])
            for offset in range(width):
                column = first_column + offset
                cell = (f"{get_column_letter(column)}{row_number}", values[offset] if offset < len(values) else None)
                cells.append(cell)
                if axis != "all":
                    buckets[row_number if axis == "row" else get_column_letter(column)].append(cell)
        sheet["summary"] = _count_cells(cells, styles, max_groups)
        if axis != "all":
            sheet["by_" + axis] = [
                {axis: position, **_count_cells(values, styles, max_groups)} for position, values in buckets.items()
            ]
    result.update(range=cell_range, axis=axis)
    return result


def _parse_cell_range(value: str) -> tuple[str | None, int, int, int, int]:
    """Convert an optional sheet name and finite A1 range to a bounded window."""
    sheet = None
    if "!" in value:
        sheet, value = value.rsplit("!", 1)
        if sheet.startswith("'") and sheet.endswith("'"):
            sheet = sheet[1:-1].replace("''", "'")
        if not sheet:
            raise ValueError("A sheet-qualified range must include a sheet name")
    first_column, first_row, last_column, last_row = range_boundaries(value)
    if any(position is None for position in (first_column, first_row, last_column, last_row)):
        raise ValueError("Select a cell or finite rectangle, such as B3 or B3:D8")
    if not (1 <= first_row <= last_row <= 1_048_576 and 1 <= first_column <= last_column <= 16_384):
        raise ValueError("Range endpoints must be ordered cells within XLSX worksheet limits")
    max_rows = last_row - first_row + 1
    max_columns = last_column - first_column + 1
    if max_rows * max_columns > 10_000:
        raise ValueError("Select at most 10,000 cells per sheet")
    return sheet, first_row, first_column, max_rows, max_columns


def main() -> None:
    """Run the workbook inspector from the sandbox shell."""
    parser = argparse.ArgumentParser(
        description=__doc__,
        epilog=(
            "Summary JSON: sheets[].summary contains whole-range counts. Optional sheets[].by_row "
            "uses worksheet row numbers, and sheets[].by_column uses column letters. "
            "Each summary has cells, blank_cells, formula_cells, groups_truncated and groups "
            "with type, value (or formula and cached_value), count and first_cell. "
            "Groups include blanks. With --styles, groups refer to the shared styles table."
        ),
    )
    parser.add_argument("path", type=Path)
    parser.add_argument("--sheet")
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument(
        "--overview",
        action="store_true",
        help="List up to 100 sheet names, visibility states, and reported sizes without reading cells.",
    )
    mode.add_argument(
        "--summary", action="store_true", help="Count typed values in an explicit --range instead of listing cells."
    )
    parser.add_argument(
        "--axis", choices=("all", "row", "column"), help="With --summary, also count each row or column. Default: all."
    )
    parser.add_argument(
        "--max-groups",
        type=int,
        help="With --summary, return up to this many groups per summary (1..1,000). Default: 100.",
    )
    parser.add_argument(
        "--styles",
        action="store_true",
        help="Include stored font/fill colors, borders, alignment, and number formats. Styles are deduplicated. Conditional formatting, tint rendering, and comments are not evaluated.",
    )
    parser.add_argument(
        "--range",
        dest="cell_range",
        help="Select at most 10,000 cells with an A1 range, such as B3:D8 or 'Sheet name'!B3:D8. Cannot be combined with --overview or row/column window options.",
    )
    parser.add_argument("--start-row", type=int)
    parser.add_argument("--start-column", type=int)
    parser.add_argument("--max-rows", type=int)
    parser.add_argument("--max-columns", type=int)
    args = parser.parse_args()
    window = (args.start_row, args.start_column, args.max_rows, args.max_columns)
    if args.summary and args.cell_range is None:
        parser.error("--summary requires --range")
    if not args.summary and (args.axis is not None or args.max_groups is not None):
        parser.error("--axis and --max-groups require --summary")
    if args.cell_range is not None:
        if args.overview or any(value is not None for value in window):
            parser.error("--range cannot be combined with --overview or row/column window options")
        try:
            sheet, *window = _parse_cell_range(args.cell_range)
        except ValueError as exc:
            parser.error(str(exc))
        if sheet is not None:
            if args.sheet is not None and args.sheet != sheet:
                parser.error("--sheet conflicts with the sheet name in --range")
            args.sheet = sheet
    else:
        window = tuple(value if value is not None else default for value, default in zip(window, (1, 1, 200, 50)))
    if args.summary:
        try:
            result = summarize_workbook(
                args.path,
                cell_range=args.cell_range,
                sheet_name=args.sheet,
                axis=args.axis or "all",
                styles=args.styles,
                max_groups=100 if args.max_groups is None else args.max_groups,
            )
        except ValueError as exc:
            parser.error(str(exc))
    else:
        result = inspect_workbook(
            args.path,
            sheet_name=args.sheet,
            start_row=window[0],
            start_column=window[1],
            max_rows=window[2],
            max_columns=window[3],
            styles=args.styles,
            overview=args.overview,
        )
    print(json.dumps(result, ensure_ascii=False, indent=None if args.overview or args.summary else 2, default=str))


if __name__ == "__main__":
    main()
