"""Anonymize optimizer submissions and restore person IDs in returned workbooks."""

# This file is part of Nurse Scheduling Project, see <https://github.com/j3soon/nurse-scheduling>.
#
# Copyright (C) 2023-2026 Johnson Sun
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as
# published by the Free Software Foundation, either version 3 of the
# License, or (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.

# This code is mostly AI generated.

from dataclasses import dataclass
from io import BytesIO
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.xml.functions import DEFUSEDXML
from ruamel.yaml import YAML

from ..loader import _load_yaml
from ..models import (
    SHIFT_AFFINITY,
    SHIFT_COUNT,
    SHIFT_REQUEST,
    SHIFT_TYPE_REQUIREMENT,
    SHIFT_TYPE_SUCCESSIONS,
)
from .validation import validate_frontend_schedule_yaml


@dataclass(frozen=True)
class PreparedOptimizerSchedule:
    """Outbound YAML and the server-only mapping needed to restore its workbook."""

    submission_yaml: str
    original_id_by_anonymized_id: dict[str, str]
    people_count: int


class OptimizerResultError(ValueError):
    """The returned workbook cannot be safely restored."""


def restore_people_ids(content: bytes, original_id_by_anonymized_id: dict[str, str], people_count: int) -> bytes:
    """Restore first-sheet person IDs before either sandbox inspection or browser download."""
    if not DEFUSEDXML:
        raise OptimizerResultError("Safe XLSX parsing is unavailable.")
    if not content.startswith(b"PK\x03\x04"):
        raise OptimizerResultError("The optimizer result is not an XLSX workbook.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 1_000 or sum(entry.file_size for entry in entries) > 50_000_000:
                raise OptimizerResultError("The optimizer workbook expands beyond the allowed size.")
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise OptimizerResultError("Encrypted optimizer workbooks are unsupported.")
            names = {entry.filename for entry in entries}
            if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
                raise OptimizerResultError("The optimizer result is not an XLSX workbook.")
    except BadZipFile as exc:
        raise OptimizerResultError("The optimizer workbook archive is invalid.") from exc
    if all(anonymized == original for anonymized, original in original_id_by_anonymized_id.items()):
        return content
    try:
        workbook = load_workbook(BytesIO(content), keep_links=False)
        try:
            if not workbook.worksheets:
                raise OptimizerResultError("The optimizer workbook has no worksheet.")
            sheet = workbook.worksheets[0]
            for row_number in range(3, 3 + people_count):
                cell = sheet.cell(row_number, 1)
                if isinstance(cell.value, str):
                    original_id = original_id_by_anonymized_id.get(cell.value)
                    if original_id is not None:
                        cell.value = original_id
            output = BytesIO()
            workbook.save(output)
            return output.getvalue()
        finally:
            workbook.close()
    except Exception as exc:
        if isinstance(exc, OptimizerResultError):
            raise
        raise OptimizerResultError("The optimizer workbook could not be restored.") from exc


def _map_references(value: Any, anonymized_id_by_original_id: dict[str, str]) -> Any:
    if isinstance(value, list):
        return [_map_references(item, anonymized_id_by_original_id) for item in value]
    if isinstance(value, str):
        return anonymized_id_by_original_id.get(value, value)
    return value


def _remove_descriptions(value: Any) -> Any:
    if isinstance(value, list):
        return [_remove_descriptions(item) for item in value]
    if isinstance(value, dict):
        return {key: _remove_descriptions(item) for key, item in value.items() if key != "description"}
    return value


def prepare_optimizer_schedule(schedule_yaml: str, max_schedule_bytes: int) -> PreparedOptimizerSchedule:
    """Replace people IDs and descriptions as the Optimize and Export page does."""
    validation = validate_frontend_schedule_yaml(schedule_yaml, max_schedule_bytes)
    if not validation.valid:
        raise ValueError(validation.render())
    payload = _load_yaml(schedule_yaml.encode("utf-8"))
    people = payload["people"]
    items = people["items"]
    # Optional sections are absent from a valid schedule that never declares them.
    groups = people.get("groups", [])
    used_ids = {group["id"] for group in groups}
    anonymized_id_by_original_id: dict[str, str] = {}
    original_id_by_anonymized_id: dict[str, str] = {}
    next_index = 1
    for item in items:
        anonymized_id = f"P{next_index}"
        while anonymized_id in used_ids:
            next_index += 1
            anonymized_id = f"P{next_index}"
        original_id = item["id"]
        anonymized_id_by_original_id[original_id] = anonymized_id
        original_id_by_anonymized_id[anonymized_id] = original_id
        used_ids.add(anonymized_id)
        item["id"] = anonymized_id
        next_index += 1

    for group in groups:
        group["members"] = _map_references(group["members"], anonymized_id_by_original_id)
    for preference in payload["preferences"]:
        kind = preference["type"]
        if kind == SHIFT_TYPE_REQUIREMENT:
            preference["qualifiedPeople"] = _map_references(
                preference.get("qualifiedPeople"), anonymized_id_by_original_id
            )
        elif kind in {SHIFT_REQUEST, SHIFT_TYPE_SUCCESSIONS, SHIFT_COUNT}:
            preference["person"] = _map_references(preference["person"], anonymized_id_by_original_id)
        elif kind == SHIFT_AFFINITY:
            for field_name in ("people1", "people2"):
                preference[field_name] = _map_references(preference[field_name], anonymized_id_by_original_id)

    export = payload.get("export") or {}
    for rule in export.get("formatting", []):
        if "people" in rule:
            rule["people"] = _map_references(rule["people"], anonymized_id_by_original_id)
    for rule in export.get("extraRows", []):
        rule["countPeople"] = _map_references(rule["countPeople"], anonymized_id_by_original_id)

    output = BytesIO()
    yaml = YAML(typ="safe")
    yaml.default_flow_style = False
    yaml.representer.sort_base_mapping_type_on_output = False
    yaml.dump(_remove_descriptions(payload), output)
    return PreparedOptimizerSchedule(output.getvalue().decode("utf-8"), original_id_by_anonymized_id, len(items))
