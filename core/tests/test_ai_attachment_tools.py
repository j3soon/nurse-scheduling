"""Tests for trusted sandbox attachment helper scripts."""

# This file is part of Nurse Scheduling Project, see <https://github.com/j3soon/nurse-scheduling>.
#
# Copyright (C) 2023-2026 Johnson Sun
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as
# published by the Free Software Foundation, either version 3 of the
# License, or (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.

# This test is mostly AI generated.

from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from nurse_scheduling.ai.attachment_tools.inspect_pdf import inspect_pdf
from nurse_scheduling.ai.attachment_tools.inspect_xlsx import inspect_workbook


def test_workbook_inspector_reads_every_sheet_including_hidden_sheets(tmp_path: Path):
    path = tmp_path / "multi-sheet.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "First"
    first.append(["person", "shift"])
    first.append(["Alice", "Day"])
    second = workbook.create_sheet("Second")
    second.append(["Bob", "Night"])
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["Private", "Value"])
    workbook.save(path)

    result = inspect_workbook(path)

    assert result["sheet_names"] == ["First", "Second", "Hidden"]
    assert [sheet["name"] for sheet in result["sheets"]] == ["First", "Second", "Hidden"]
    assert result["sheets"][1]["rows"] == [{"row": 1, "values": ["Bob", "Night"]}]
    assert result["sheets"][2]["state"] == "hidden"


def test_workbook_inspector_shows_formulas_and_last_saved_values(tmp_path: Path):
    path = tmp_path / "formula.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = 2
    worksheet["A2"] = 3
    worksheet["A3"] = "=SUM(A1:A2)"
    source = BytesIO()
    workbook.save(source)
    workbook.close()
    source.seek(0)
    with ZipFile(source) as archive, ZipFile(path, "w", ZIP_DEFLATED) as updated:
        for entry in archive.infolist():
            content = archive.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b"<f>SUM(A1:A2)</f><v />", b"<f>SUM(A1:A2)</f><v>5</v>")
            updated.writestr(entry, content)

    result = inspect_workbook(path)

    assert result["cached_values_are_last_saved_not_recalculated"]
    assert result["sheets"][0]["rows"][2] == {
        "row": 3,
        "values": [{"cell": "A3", "formula": "=SUM(A1:A2)", "cached_value": 5}],
    }


def test_workbook_inspector_selects_later_rows_and_sheets(tmp_path: Path):
    path = tmp_path / "large.xlsx"
    workbook = Workbook()
    workbook.active["A201"] = "Later row"
    for index in range(21):
        workbook.create_sheet(f"Sheet {index}")
    workbook["Sheet 20"]["C4"] = "Later sheet"
    workbook.save(path)

    overview = inspect_workbook(path)
    later_row = inspect_workbook(path, start_row=201, max_rows=1, max_columns=1)
    later_sheet = inspect_workbook(path, sheet_name="Sheet 20", start_row=4, start_column=3, max_rows=1, max_columns=1)

    assert overview["sheets_truncated"]
    assert overview["sheets"][0]["truncated"]
    assert later_row["sheets"][0]["rows"] == [{"row": 201, "values": ["Later row"]}]
    assert later_sheet["sheets"][0]["rows"] == [{"row": 4, "values": ["Later sheet"]}]
    assert later_sheet["sheets"][0]["first_column"] == 3


def _text_pdf(path: Path) -> None:
    writer = PdfWriter()
    for label in ("First page", "Second page"):
        page = writer.add_blank_page(width=300, height=200)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 18 Tf 40 100 Td ({label}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)


def test_pdf_inspector_extracts_text_by_page_and_reports_truncation(tmp_path: Path):
    path = tmp_path / "text.pdf"
    _text_pdf(path)

    first = inspect_pdf(path, max_pages=1)
    second = inspect_pdf(path, page_number=2, render=True, output_directory=tmp_path / "rendered")

    assert first["page_count"] == 2
    assert first["pages_truncated"]
    assert "First page" in first["pages"][0]["text"]
    assert second["pages"][0]["page"] == 2
    assert "Second page" in second["pages"][0]["text"]
    assert not second["ocr_performed"]
    with Image.open(second["pages"][0]["rendered_image"]["path"]) as image:
        assert image.convert("L").getextrema()[0] < 255

    shortened = inspect_pdf(path, page_number=2, max_chars_per_page=6)
    assert shortened["pages"][0]["text"] == "Second"
    assert shortened["pages"][0]["text_truncated"]


def test_pdf_inspector_renders_an_image_only_page(tmp_path: Path):
    path = tmp_path / "image.pdf"
    Image.new("RGB", (100, 80), "blue").save(path, "PDF")

    result = inspect_pdf(path, page_number=1, render=True, output_directory=tmp_path / "rendered")

    page = result["pages"][0]
    assert not page["has_extractable_text"]
    image_path = Path(page["rendered_image"]["path"])
    with Image.open(image_path) as image:
        assert image.format == "PNG"
        assert image.width > 0 and image.height > 0


def test_pdf_inspector_requires_a_selected_page_for_rendering(tmp_path: Path):
    path = tmp_path / "image.pdf"
    Image.new("RGB", (10, 10), "red").save(path, "PDF")

    with pytest.raises(ValueError, match="Select one page"):
        inspect_pdf(path, render=True)
