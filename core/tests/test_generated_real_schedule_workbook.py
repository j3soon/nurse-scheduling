"""Check that the bundled ward workbook matches its generator."""

# This file is part of Nurse Scheduling Project, see <https://github.com/j3soon/nurse-scheduling>.
#
# Copyright (C) 2023-2026 Johnson Sun
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as
# published by the Free Software Foundation, either version 3 of the
# License, or (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.

# This test is mostly AI generated.

import subprocess
import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
GENERATOR = ROOT / "scripts/generate_real_schedule_workbook.py"
BUNDLED = ROOT / "docs/content/user-guide/build-a-real-schedule/unfilled-ward-schedule-2025-11.xlsx"


def _cell_content(cell):
    content = (cell.value, cell.data_type, cell.has_style)
    if not cell.has_style:
        return content
    return content + (
        cell.number_format,
        copy(cell.font),
        copy(cell.fill),
        copy(cell.border),
        copy(cell.alignment),
        copy(cell.protection),
    )


def _column_layout(sheet):
    return {
        key: (item.width, item.hidden, item.min, item.max, item.collapsed, item.style)
        for key, item in sheet.column_dimensions.items()
    }


def test_bundled_real_schedule_workbook_matches_generator(tmp_path):
    generated = tmp_path / "generated.xlsx"
    subprocess.run(
        [sys.executable, str(GENERATOR), "--output", str(generated)],
        cwd=tmp_path,
        check=True,
        capture_output=True,
        text=True,
    )

    bundled_book = load_workbook(BUNDLED)
    generated_book = load_workbook(generated)
    try:
        assert generated_book.sheetnames == bundled_book.sheetnames
        for field in ("title", "subject", "creator", "keywords"):
            assert getattr(generated_book.properties, field) == getattr(bundled_book.properties, field)

        for name in bundled_book.sheetnames:
            bundled_sheet = bundled_book[name]
            generated_sheet = generated_book[name]
            assert generated_sheet.sheet_state == bundled_sheet.sheet_state
            assert (generated_sheet.max_row, generated_sheet.max_column) == (
                bundled_sheet.max_row,
                bundled_sheet.max_column,
            )
            assert generated_sheet.freeze_panes == bundled_sheet.freeze_panes
            assert generated_sheet.print_title_rows == bundled_sheet.print_title_rows
            assert _column_layout(generated_sheet) == _column_layout(bundled_sheet)
            assert set(generated_sheet.merged_cells.ranges) == set(bundled_sheet.merged_cells.ranges)

            for row in bundled_sheet:
                for bundled_cell in row:
                    generated_cell = generated_sheet[bundled_cell.coordinate]
                    assert _cell_content(generated_cell) == _cell_content(bundled_cell), (
                        f"{name}!{bundled_cell.coordinate} differs"
                    )
    finally:
        bundled_book.close()
        generated_book.close()
