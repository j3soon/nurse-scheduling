"""Semantic and standalone checks for the sandbox optimizer result reader."""

# This file is part of Nurse Scheduling Project, see <https://github.com/j3soon/nurse-scheduling>.
#
# Copyright (C) 2023-2026 Johnson Sun
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as
# published by the Free Software Foundation, either version 3 of the
# License, or (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.

# This test is mostly AI generated.

import json
import subprocess
import sys

import pytest
from openpyxl import load_workbook

from nurse_scheduling.ai.attachment_tools.inspect_optimizer_result import (
    _assigned,
    inspect_result,
    load_context,
    summarize_roster,
)
from nurse_scheduling.ai.result_context import MAX_REQUEST_AUDIT_BYTES, build_request_audit, build_result_context
from nurse_scheduling.ai.workspace import INSPECTION_HELPERS, REFERENCE_ATTACHMENT_TOOLS, inspection_helper_catalog

from .ai_eval.optimizer_fixtures import FIXTURE, RESULT_SOURCES, completion_result
from .ai_test_helper import parse_schedule, schedule_yaml


@pytest.fixture
def audit(tmp_path):
    context = build_result_context(FIXTURE.read_text())
    payload, _ = completion_result("request-audit", FIXTURE.read_text())
    path = tmp_path / "result.xlsx"
    path.write_bytes(payload)
    return path, context


def test_reader_counts_actual_assignments_and_hard_avoids(audit):
    path, context = audit
    result = inspect_result(path, context, context["source_sha256"])
    assert result["summary"] == [
        {"weight": 11_000_000_000, "total": 4, "satisfied": 3, "unmet": 1},
        {"weight": 11_000_000, "total": 1, "satisfied": 0, "unmet": 1},
        {"weight": "-.inf", "total": 3, "satisfied": 3, "unmet": 0},
    ]
    assert [(row["person"], row["date"], row["assigned"]) for row in result["unmet"]] == [
        ("Mira", "2026-05-02", "N"),
        ("Alex", "2026-05-03", "D"),
    ]
    assert result["score"] == 33_000_000_000
    assert result["score_direction"] == "maximize"
    assert result["status"] == "FEASIBLE"


def test_completion_summary_has_bounded_counts_and_explicit_scope(audit):
    path, context = audit
    result = build_request_audit(FIXTURE.read_text(), path.read_bytes())
    assert result is not None
    assert result["source_sha256"] == context["source_sha256"]
    assert result["summary"] == [
        {"weight": 11_000_000_000, "total": 4, "satisfied": 3, "unmet": 1},
        {"weight": 11_000_000, "total": 1, "satisfied": 0, "unmet": 1},
        {"weight": "-.inf", "total": 3, "satisfied": 3, "unmet": 0},
    ]
    assert "Staffing and rest are not audited" in result["scope"]
    assert set(result) == {"scope", "source_sha256", "summary"}
    assert len(json.dumps(result).encode()) <= MAX_REQUEST_AUDIT_BYTES


def test_summary_reports_when_weight_tiers_exceed_limit(audit):
    path, _ = audit
    payload = parse_schedule(FIXTURE.read_text())
    request = next(pref for pref in payload["preferences"] if pref["type"] == "shift request")
    required = [pref for pref in payload["preferences"] if pref["type"] != "shift request"]
    payload["preferences"] = required + [{**request, "weight": weight} for weight in range(1, 101)]
    assert build_request_audit(schedule_yaml(payload), path.read_bytes()) == {
        "unavailable": "The automatic request audit exceeds the summary size limit."
    }


def test_summary_falls_back_for_invalid_and_unsupported_workbooks(audit):
    path, _ = audit
    assert "could not be read" in build_request_audit(FIXTURE.read_text(), b"not a workbook")["unavailable"]
    workbook = load_workbook(path)
    workbook.active.cell(1, 5).value = "unsupported date header"
    workbook.save(path)
    assert "could not be read" in build_request_audit(FIXTURE.read_text(), path.read_bytes())["unavailable"]


def test_stale_summary_control_disagrees_with_current_verified_incumbent():
    source = FIXTURE.read_text()
    workbook, metadata = completion_result("request-audit-stale-summary", source)
    current = build_request_audit(source, workbook)
    assert current is not None
    assert metadata["request_audit"]["source_sha256"] != metadata["source_sha256"]
    assert current["source_sha256"] == metadata["source_sha256"]
    assert metadata["request_audit"]["summary"][0]["unmet"] == 1
    assert current["summary"][0]["unmet"] == 0


def test_reader_accepts_reordered_people_and_offset_headers(audit):
    path, context = audit
    workbook = load_workbook(path)
    ws = workbook.active
    first, third = [c.value for c in ws[3]], [c.value for c in ws[5]]
    for col, value in enumerate(third, 1):
        ws.cell(3, col).value = value
    for col, value in enumerate(first, 1):
        ws.cell(5, col).value = value
    ws.insert_rows(1)
    ws.cell(1, 1).value = "Roster title"
    workbook.save(path)
    result = inspect_result(path, context, context["source_sha256"])
    assert result["summary"][0]["unmet"] == 1


def test_reader_does_not_count_lower_weight_markers_as_strong_misses(audit):
    path, context = audit
    payload, _ = completion_result("request-audit-all-strong", FIXTURE.read_text())
    path.write_bytes(payload)
    result = inspect_result(path, context, context["source_sha256"])
    assert [(row["total"], row["unmet"]) for row in result["summary"]] == [(4, 0), (1, 1), (3, 0)]


def test_compiled_context_uses_compiled_groups_and_reserved_selectors():
    source = (
        FIXTURE.read_text()
        .replace("groups: []\nshiftTypes:", "groups: [{id: Team, members: [Kai]}]\nshiftTypes:")
        .replace("person: [Kai]", "person: [Team]")
    )
    context = build_result_context(source)
    last = context["requests"][-1]
    assert last["people"] == ["Kai"]
    assert last["dates"] == ["2026-05-01", "2026-05-02", "2026-05-03"]
    assert last["shift_types"] == ["K"]


def test_reader_audits_nested_groups_cross_month_dates_and_negative_weights(tmp_path):
    source = RESULT_SOURCES["request-audit-groups"].read_text()
    context = build_result_context(source)
    payload, _ = completion_result("request-audit-groups", source)
    path = tmp_path / "groups.xlsx"
    path.write_bytes(payload)
    result = inspect_result(path, context, context["source_sha256"])
    assert [(row["weight"], row["total"], row["unmet"]) for row in result["summary"]] == [
        (11_000_000_000, 8, 3),
        (11_000_000, 6, 2),
        (-11_000_000_000, 6, 2),
        ("-.inf", 24, 0),
    ]
    assert context["requests"][0]["people"] == ["Asha", "Ben", "Cleo"]
    assert context["requests"][3]["shift_types"] == ["D", "N"]


def test_reader_rejects_mismatched_source_and_unknown_decorations(audit):
    path, context = audit
    with pytest.raises(ValueError, match="source"):
        inspect_result(path, context, "0" * 64)
    workbook = load_workbook(path)
    workbook.active.cell(3, 5).value = "* OFF?"
    workbook.save(path)
    with pytest.raises(ValueError, match="unknown or ambiguous"):
        inspect_result(path, context, context["source_sha256"])


@pytest.mark.parametrize(
    "value,expected", [(None, "OFF"), (" [OFF]", "OFF"), ("K [K]", "K"), ("D [OFF] [X]", "D"), ("D+", "D+"), (0, "0")]
)
def test_assignment_reader_handles_annotations_and_exact_shift_ids(value, expected):
    assert _assigned(value, ["D", "D+", "K", "0"]) == expected


def test_assignment_reader_rejects_ambiguous_literal_shift_ids():
    with pytest.raises(ValueError, match="ambiguous"):
        _assigned("D [OFF]", ["D", "D [OFF]"])


def test_reader_bounds_detail_without_changing_counts(audit):
    path, context = audit
    result = inspect_result(path, context, context["source_sha256"], max_unmet=1)
    assert len(result["unmet"]) == 1 and result["unmet_truncated"]
    filtered = inspect_result(path, context, context["source_sha256"], weights=["11000000000"])
    assert len(filtered["summary"]) == 1 and filtered["summary"][0]["total"] == 4


def test_assignment_query_uses_actual_shifts_and_reports_truncation(audit):
    path, context = audit
    result = inspect_result(
        path,
        context,
        context["source_sha256"],
        people=["Mira", "Alex"],
        dates=["2026-05-02", "2026-05-03"],
        max_assignments=3,
    )
    assert result["assignments"] == [
        {"person": "Mira", "date": "2026-05-02", "shift_type": "N"},
        {"person": "Mira", "date": "2026-05-03", "shift_type": "OFF"},
        {"person": "Alex", "date": "2026-05-02", "shift_type": "K"},
    ]
    assert result["assignments_truncated"]
    assert "summary" not in result


@pytest.mark.parametrize(
    "filters", [{"people": ["missing"]}, {"dates": ["02"]}, {"people": ["Alex"], "max_assignments": 0}]
)
def test_assignment_query_rejects_unknown_selectors_and_unbounded_limits(audit, filters):
    path, context = audit
    with pytest.raises(ValueError):
        inspect_result(path, context, context["source_sha256"], **filters)


@pytest.mark.parametrize(
    "fixture,cli_args,expected",
    [
        ("request-audit", [], [(11_000_000_000, 4, 1), (11_000_000, 1, 1), ("-.inf", 3, 0)]),
        ("request-audit", ["--weight", "-.inf"], [("-.inf", 3, 0)]),
        ("request-audit", ["--weight=-.inf"], [("-.inf", 3, 0)]),
        ("request-audit-groups", ["--weight", "-11e9"], [(-11_000_000_000, 6, 2)]),
        (
            "request-audit",
            ["--person", "Mira", "--date", "2026-05-02"],
            [
                {"person": "Mira", "date": "2026-05-02", "shift_type": "N"},
            ],
        ),
        (
            "request-audit-groups",
            ["--person", "Asha", "--date", "2026-06-01"],
            [
                {"person": "Asha", "date": "2026-06-01", "shift_type": "N"},
            ],
        ),
    ],
)
def test_standalone_hydrated_reader_runs_without_project_imports(tmp_path, fixture, cli_args, expected):
    source = RESULT_SOURCES[fixture].read_text()
    context = build_result_context(source)
    path = tmp_path / "result.xlsx"
    path.write_bytes(completion_result(fixture, source)[0])
    for source in REFERENCE_ATTACHMENT_TOOLS.values():
        (tmp_path / source.name).write_bytes(source.read_bytes())
    context_path = tmp_path / "context.json"
    context_path.write_text(json.dumps(context))
    result = subprocess.run(
        [
            sys.executable,
            str(tmp_path / "inspect_optimizer_result.py"),
            str(path),
            "--context",
            str(context_path),
            "--source-sha256",
            context["source_sha256"],
            *cli_args,
        ],
        cwd=tmp_path,
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    parsed = json.loads(result.stdout)
    if "--person" in cli_args:
        assert parsed["assignments"] == expected
        assert not parsed["assignments_truncated"]
    else:
        assert [(row["weight"], row["total"], row["unmet"]) for row in parsed["summary"]] == expected


def test_catalog_only_advertises_real_hydrated_scripts():
    catalog = inspection_helper_catalog()
    assert len(INSPECTION_HELPERS) == len(REFERENCE_ATTACHMENT_TOOLS)
    for path, source in REFERENCE_ATTACHMENT_TOOLS.items():
        assert source.is_file() and path in catalog


def test_default_context_must_match_the_completion_source(tmp_path, monkeypatch):
    from nurse_scheduling.ai.attachment_tools import inspect_optimizer_result as reader

    approved, pending = tmp_path / "schedule-context.json", tmp_path / "pending-schedule-context.json"
    approved.write_text(json.dumps({"schema_version": 1, "source_sha256": "approved"}))
    pending.write_text(json.dumps({"schema_version": 1, "source_sha256": "pending"}))
    monkeypatch.setattr(reader, "CONTEXT", str(approved))
    assert load_context("approved")["source_sha256"] == "approved"
    assert load_context("pending")["source_sha256"] == "pending"
    with pytest.raises(ValueError, match="No hydrated"):
        load_context("unknown")
    approved.write_text("broken JSON")
    assert load_context("pending")["source_sha256"] == "pending"
    pending.write_text(json.dumps({"schema_version": 2, "source_sha256": "pending"}))
    with pytest.raises(ValueError, match="No hydrated"):
        load_context("pending")


@pytest.mark.parametrize("explicit_stale", [False, True])
def test_result_cli_selects_pending_source_and_respects_explicit_context(tmp_path, monkeypatch, capsys, explicit_stale):
    from nurse_scheduling.ai.attachment_tools import inspect_optimizer_result as reader

    approved, pending = tmp_path / "schedule-context.json", tmp_path / "pending-schedule-context.json"
    approved_context = build_result_context(FIXTURE.read_text())
    pending_source = RESULT_SOURCES["request-audit-pending"].read_text()
    pending_context = build_result_context(pending_source)
    approved.write_text(json.dumps(approved_context))
    pending.write_text(json.dumps(pending_context))
    workbook, _ = completion_result("request-audit-pending", pending_source)
    path = tmp_path / "result.xlsx"
    path.write_bytes(workbook)
    monkeypatch.setattr(reader, "CONTEXT", str(approved))
    arguments = [
        "inspect_optimizer_result.py",
        str(path),
        "--source-sha256",
        pending_context["source_sha256"],
        "--weight",
        "11000000000",
    ]
    if explicit_stale:
        arguments += ["--context", str(approved)]
    monkeypatch.setattr(sys, "argv", arguments)
    if explicit_stale:
        with pytest.raises(SystemExit) as error:
            reader.main()
        assert error.value.code == 1
        assert "does not match" in capsys.readouterr().err
    else:
        with pytest.raises(ValueError, match="does not match"):
            inspect_result(path, approved_context, pending_context["source_sha256"])
        reader.main()
        result = json.loads(capsys.readouterr().out)
        assert result["summary"] == [{"weight": 11000000000, "total": 5, "satisfied": 3, "unmet": 2}]


def test_roster_statistics_count_exact_shifts_and_bound_people(audit):
    path, context = audit
    result = summarize_roster(path, context, context["source_sha256"], max_people=2)
    assert result["statistics_truncated"]
    assert result["statistics"][0] == {
        "person": "Alex",
        "shift_counts": {"OFF": 1, "K": 1, "D": 1},
        "max_work_run": 2,
        "max_off_run": 1,
    }
    assert result["statistics"][1]["person"] == "Mira"
    assert "not a policy audit" in result["scope"]
    assert result["summary"]["people"] == 3
    assert sum(result["summary"]["shift_counts"].values()) == 9


def test_roster_statistics_include_only_history_runs_crossing_the_boundary(audit, tmp_path):
    path, _ = audit
    payload = parse_schedule(FIXTURE.read_text(encoding="utf-8"))
    for person, history in zip(payload["people"]["items"], [["OFF", "OFF"], ["D", "N"], None]):
        person["history"] = history
    source = tmp_path / "source.yaml"
    source.write_bytes(schedule_yaml(payload).encode())
    context = build_result_context(source.read_text(encoding="utf-8"))
    result = summarize_roster(path, context, context["source_sha256"], include_history=True, source=source)
    alex, mira, kai = result["statistics"]
    assert alex["max_off_run_including_history"] == 3
    assert alex["shift_counts"]["OFF"] == 1
    assert mira["max_work_run_including_history"] == 4
    assert kai["max_work_run_including_history"] is None
    assert kai["history_days"] is None
    # Changed source history must not be used with an earlier optimizer result.
    source.write_bytes(source.read_bytes() + b"\n# changed")
    alex = summarize_roster(path, context, context["source_sha256"], include_history=True, source=source)["statistics"][
        0
    ]
    assert alex["history_days"] is None


@pytest.mark.parametrize("matching_source", [False, True], ids=["different-bytes", "matching-bytes"])
def test_roster_history_checks_exact_source_bytes(audit, tmp_path, matching_source):
    path, _ = audit
    text = FIXTURE.read_text(encoding="utf-8")
    windows_text = text.replace("\n", "\r\n")
    source = tmp_path / "source.yaml"
    source.write_bytes(windows_text.encode())
    context = build_result_context(windows_text if matching_source else text)
    row = summarize_roster(path, context, context["source_sha256"], source=source, include_history=True)["statistics"][
        0
    ]
    assert row["history_days"] == (3 if matching_source else None)


@pytest.mark.parametrize("options", [{"max_people": 0}, {"max_people": 1001}, {"people": ["missing"]}])
def test_roster_statistics_reject_unknown_people_and_unbounded_options(audit, options):
    path, context = audit
    with pytest.raises(ValueError):
        summarize_roster(path, context, context["source_sha256"], **options)


def test_statistics_preserve_qualified_shifts_and_ignore_annotations(audit):
    path, context = audit
    workbook = load_workbook(path)
    context["shift_types"].append("D+")
    for row in workbook.active:
        for cell in row:
            if str(cell.value) == "D" or str(cell.value).startswith("D ["):
                cell.value = "D+ [OFF]"
    workbook.save(path)
    result = summarize_roster(path, context, context["source_sha256"])
    alex = result["statistics"][0]
    assert alex["shift_counts"] == {"OFF": 1, "K": 1, "D+": 1}
    assert alex["max_work_run"] == 2


@pytest.mark.parametrize("summary_only", [False, True])
def test_standalone_roster_statistics_reuse_the_hydrated_readers(audit, tmp_path, summary_only):
    path, context = audit
    for source in REFERENCE_ATTACHMENT_TOOLS.values():
        (tmp_path / source.name).write_bytes(source.read_bytes())
    source = tmp_path / "source.yaml"
    source.write_bytes(FIXTURE.read_text(encoding="utf-8").encode())
    context_path = tmp_path / "context.json"
    context_path.write_text(json.dumps(context), encoding="utf-8")
    result = subprocess.run(
        [
            sys.executable,
            str(tmp_path / "inspect_optimizer_result.py"),
            "--statistics",
            *(["--summary-only"] if summary_only else []),
            str(path),
            "--context",
            str(context_path),
            "--source-sha256",
            context["source_sha256"],
            "--source",
            str(source),
            "--include-history",
            "--person",
            "Alex",
        ],
        cwd=tmp_path,
        capture_output=True,
        text=True,
        check=True,
    )
    output = json.loads(result.stdout)
    assert output["summary"]["shift_counts"] == {"OFF": 1, "K": 1, "D": 1}
    assert output["summary"]["max_work_run_including_history"] == 2
    if summary_only:
        assert output["statistics"] == []
        return
    row = output["statistics"][0]
    assert row["history_days"] == 3
    assert row["shift_counts"] == {"OFF": 1, "K": 1, "D": 1}
    assert row["max_work_run_including_history"] == 2


def test_history_only_streak_does_not_inflate_scheduled_work_run(audit, tmp_path):
    path, _ = audit
    payload = parse_schedule(FIXTURE.read_text(encoding="utf-8"))
    payload["people"]["items"][0]["history"] = ["N"] * 10
    source = tmp_path / "source.yaml"
    source.write_bytes(schedule_yaml(payload).encode())
    context = build_result_context(source.read_text(encoding="utf-8"))
    row = summarize_roster(path, context, context["source_sha256"], source=source, include_history=True)["statistics"][
        0
    ]
    assert row["history_days"] == 10
    assert row["max_work_run_including_history"] == 2


@pytest.mark.parametrize(
    "options",
    [
        ["--summary-only"],
        ["--max-people", "1"],
        ["--include-history"],
        ["--source", "source.yaml"],
        ["--statistics", "--weight", "11"],
        ["--statistics", "--max-unmet", "20"],
        ["--statistics", "--date", "2026-05-01"],
        ["--statistics", "--max-assignments", "100"],
    ],
)
def test_result_cli_rejects_options_from_another_mode(monkeypatch, capsys, options):
    from nurse_scheduling.ai.attachment_tools.inspect_optimizer_result import main

    monkeypatch.setattr(sys, "argv", ["inspect_optimizer_result.py", "--source-sha256", "unused", *options])
    with pytest.raises(SystemExit) as error:
        main()
    assert error.value.code == 2
    assert "--statistics" in capsys.readouterr().err


def test_monthly_statistics_match_independent_counts_and_run_lengths(tmp_path):
    source = RESULT_SOURCES["monthly-roster-statistics"].read_text(encoding="utf-8")
    context = build_result_context(source)
    path = tmp_path / "monthly.xlsx"
    path.write_bytes(completion_result("monthly-roster-statistics", source)[0])
    rows = summarize_roster(path, context, context["source_sha256"])["statistics"]
    summary = summarize_roster(path, context, context["source_sha256"], summary_only=True)["summary"]
    assert summary == {
        "people": 60,
        "shift_counts": {"D": 232, "D+": 232, "E": 465, "N": 466, "OFF": 465},
        "max_work_run": 6,
        "max_work_run_distribution": {6: 60},
        "max_off_run": 2,
        "max_off_run_distribution": {2: 60},
    }
    assert len(rows) == 60
    assert {
        shift: sum(row["shift_counts"].get(shift, 0) for row in rows) for shift in ["D", "D+", "E", "N", "OFF"]
    } == {"D": 232, "D+": 232, "E": 465, "N": 466, "OFF": 465}
    assert all(row["max_work_run"] == 6 and row["max_off_run"] == 2 for row in rows)


def test_roster_summary_selects_people_before_limiting_rows(audit):
    path, context = audit
    full = summarize_roster(path, context, context["source_sha256"])
    bounded = summarize_roster(path, context, context["source_sha256"], max_people=1)
    assert bounded["summary"] == full["summary"]
    selected = summarize_roster(path, context, context["source_sha256"], people=["Alex", "Alex"], summary_only=True)
    assert selected["statistics"] == []
    assert selected["statistics_truncated"]
    assert selected["summary"] == {
        "people": 1,
        "shift_counts": {"OFF": 1, "K": 1, "D": 1},
        "max_work_run": 2,
        "max_work_run_distribution": {2: 1},
        "max_off_run": 1,
        "max_off_run_distribution": {1: 1},
    }


def test_roster_summary_keeps_unknown_history_separate(audit):
    path, context = audit
    summary = summarize_roster(path, context, context["source_sha256"], include_history=True)["summary"]
    assert summary["max_work_run_including_history"] is None
    assert summary["max_work_run_including_history_distribution"] == {}
    assert summary["max_work_run_including_history_unknown_people"] == 3
    assert sum(summary["max_work_run_distribution"].values()) == 3
