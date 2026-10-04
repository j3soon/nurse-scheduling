"""Tests for the evaluation runner, using a scripted provider."""

# This file is part of Nurse Scheduling Project, see <https://github.com/j3soon/nurse-scheduling>.
#
# Copyright (C) 2023-2026 Johnson Sun
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as
# published by the Free Software Foundation, either version 3 of the
# License, or (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.

# This test is mostly AI generated.

import asyncio
import hashlib
import json
import subprocess
import sys
from collections.abc import AsyncIterator, Sequence
from copy import deepcopy
from dataclasses import replace
from pathlib import Path

import pytest
from ruamel.yaml import YAML

from nurse_scheduling.ai.config import AiSettings
from nurse_scheduling.ai.optimizer import optimizer_start_message
from nurse_scheduling.ai.pi.bash import BASH_TOOL
from nurse_scheduling.ai.pi.edit import EDIT_TOOL
from nurse_scheduling.ai.pi.read import READ_TOOL
from nurse_scheduling.ai.provider import (
    ChatMessage,
    ProviderAttempt,
    ProviderError,
    ProviderResponseLimitError,
    ReasoningDelta,
    TextDelta,
    TokenUsage,
    ToolCall,
    ToolCallRequest,
)
from nurse_scheduling.ai.sandbox import CommandResult, SandboxError
from nurse_scheduling.ai.sandbox.fake import FakeSandboxBackend, FakeSandboxFactory
from nurse_scheduling.ai.sandbox_agent import (
    REFERENCE_ATTACHMENT_TOOLS,
    SANDBOX_SYSTEM_PROMPT,
    WORKSPACE_SCHEDULE,
    SandboxTurnMetrics,
)
from nurse_scheduling.ai.schema import (
    SCHEMA_REFERENCE_FILES,
    TAIWAN_HOLIDAYS_SOURCE,
    load_user_guide_references,
)
from nurse_scheduling.ai.system_prompt import PROMPT_DIRECTORY, compose_system_prompt, load_system_prompt_sections
from nurse_scheduling.ai.validation import validate_frontend_schedule_yaml
from nurse_scheduling.loader import _load_yaml

from .ai_eval.comparison import comparison_metrics_markdown, comparison_statistics
from .ai_eval.grading import EvalCase, ExpectedDiff, RunOutcome, ToolUsageExpectation, TurnAction, grade, load_cases
from .ai_eval.prompt_ladder import STEPS_PATH, case_digest, load_prompt_steps, prompt_at_step, validate_prompt_evidence
from .ai_eval.runner import (
    CASES,
    DEFAULT_CASE_JOBS,
    FIXTURES,
    CaseRun,
    _evaluation_metadata,
    _parse_args,
    _reference_digests,
    _selected_cases,
    default_output_dir,
    fixture_text,
    main,
    prompt_comparison_markdown,
    run_all,
    run_case,
    select,
    summarize,
    write_report,
)

CASE_BY_ID = {case.id: case for case in load_cases(CASES)}
FIXTURE_DIGESTS = {
    fixture: hashlib.sha256(fixture_text(fixture).encode()).hexdigest()
    for fixture in {case.fixture for case in CASE_BY_ID.values()}
}


def test_ai_eval_defaults_to_four_concurrent_cases():
    assert DEFAULT_CASE_JOBS == 4


def test_offline_yaml_clause_can_be_omitted_for_comparison():
    steps = load_prompt_steps()
    index = next(index for index, step in enumerate(steps, 1) if step.id == "offline-yaml")
    before = compose_system_prompt(omit=index)
    after = compose_system_prompt()
    assert before == prompt_at_step(len(steps), omit=index)
    assert after == SANDBOX_SYSTEM_PROMPT
    clause = load_system_prompt_sections()[index - 1]
    assert clause not in before
    assert after.count(clause) == 1


def test_prompt_steps_reconstruct_production_and_link_real_cases():
    steps = load_prompt_steps()
    assert steps
    assert {step.file for step in steps} == {f"steps/{path.name}" for path in (PROMPT_DIRECTORY / "steps").glob("*.md")}
    assert prompt_at_step(len(steps)) == SANDBOX_SYSTEM_PROMPT
    assert not prompt_at_step(0)
    assert all(case_id in CASE_BY_ID for step in steps for case_id in step.cases)
    sections = load_system_prompt_sections()
    for index in range(1, len(steps) + 1):
        assert prompt_at_step(index).startswith(prompt_at_step(index - 1))
        assert sections[index - 1] not in prompt_at_step(index, omit=index)


def test_every_shipped_prompt_clause_has_current_repeated_benefit_evidence():
    validate_prompt_evidence(load_prompt_steps(), CASE_BY_ID, FIXTURE_DIGESTS)


def test_optimizer_goal_and_background_start_have_separate_ladder_comparisons():
    steps = load_prompt_steps()
    sections = load_system_prompt_sections()
    expected = {
        "optimizer-goal": {
            "optimizer-goal-preserves-policy",
            "optimizer-authorized-policy-change",
            "optimizer-independent-request-check",
        },
        "optimizer-start": {"tool-optimizer-start-without-chat-hint"},
    }
    for identifier, targeted_cases in expected.items():
        index = next(index for index, step in enumerate(steps, 1) if step.id == identifier)
        arguments, cases = _selected_cases(["--prompt-compare-step", str(index)])
        assert arguments.repeat == 3 and arguments.jobs == 4
        assert steps[index - 1].comparison_mode == "adjacent"
        assert {case.id for case in cases} == targeted_cases
        assert prompt_at_step(index) == prompt_at_step(index - 1) + "\n\n" + sections[index - 1]
    identifiers = [step.id for step in steps]
    assert identifiers.index("optimizer-start") < identifiers.index("optimizer-goal")


@pytest.mark.parametrize(
    ("change", "message"),
    [
        ({"input_sha256": "0" * 64}, "stale clause, testcase, or fixture inputs"),
        ({"before": "3/3", "after": "3/3"}, "no measured benefit"),
        ({"infrastructure_errors": 1}, "infrastructure errors"),
        ({"after": "2/3"}, "all after attempts passing"),
        ({"before": "1/1", "after": "1/1"}, "three to ten paired runs"),
    ],
)
def test_prompt_evidence_rejects_stale_or_inconclusive_receipts(change, message):
    step = load_prompt_steps()[0]
    record = {**step.evidence[0], **change}
    with pytest.raises(ValueError, match=message):
        validate_prompt_evidence((replace(step, evidence=(record,)),), CASE_BY_ID, FIXTURE_DIGESTS)


def test_prompt_evidence_rejects_a_regressing_control():
    step = next(step for step in load_prompt_steps() if step.evidence[0].get("controls"))
    record = {**step.evidence[0], "controls": [{**step.evidence[0]["controls"][0], "after": "2/3"}]}
    with pytest.raises(ValueError, match="all after attempts passing"):
        validate_prompt_evidence((replace(step, evidence=(record,)),), CASE_BY_ID, FIXTURE_DIGESTS)


def test_prompt_evidence_requires_a_witness():
    with pytest.raises(ValueError, match="no benefit witness"):
        validate_prompt_evidence((replace(load_prompt_steps()[0], evidence=()),), CASE_BY_ID, FIXTURE_DIGESTS)


def test_attachment_evidence_binds_reproducible_binary_inputs(monkeypatch):
    from .ai_eval import prompt_ladder

    case = CASE_BY_ID["xlsx-formulas-and-caches"]
    digest = case_digest(case)
    assert digest == case_digest(case)
    original = prompt_ladder.load_attachment_fixtures
    monkeypatch.setattr(
        prompt_ladder,
        "load_attachment_fixtures",
        lambda names: tuple(replace(attachment, data=attachment.data + b"changed") for attachment in original(names)),
    )
    assert case_digest(case) != digest


@pytest.mark.parametrize("changed_input", ["clause", "case", "fixture"])
def test_prompt_evidence_rejects_changed_inputs(changed_input):
    step = load_prompt_steps()[0]
    case = CASE_BY_ID[step.evidence[0]["case"]]
    cases, fixtures = CASE_BY_ID, FIXTURE_DIGESTS
    if changed_input == "clause":
        step = replace(step, sha256="0" * 64)
    elif changed_input == "case":
        cases = {**cases, case.id: replace(case, user_turns=("A changed request.",))}
    else:
        fixtures = {**fixtures, case.fixture: "0" * 64}
    with pytest.raises(ValueError, match="stale clause, testcase, or fixture inputs"):
        validate_prompt_evidence((step,), cases, fixtures)


@pytest.mark.parametrize("metric", ["tool-calls", "completion-tokens", "total-tokens"])
@pytest.mark.parametrize("ratio, passes", [(0.55, True), (0.70, False)])
def test_prompt_evidence_accepts_only_cost_gains_meeting_the_declared_target(ratio, passes, metric):
    step = load_prompt_steps()[0]
    record = {
        **step.evidence[0],
        "before": "3/3",
        "after": "3/3",
        "cost_metric": metric,
        "cost_ratio": ratio,
        "cost_target": 0.60,
    }
    candidate = (replace(step, evidence=(record,)),)
    if passes:
        validate_prompt_evidence(candidate, CASE_BY_ID, FIXTURE_DIGESTS)
    else:
        with pytest.raises(ValueError, match="no measured benefit"):
            validate_prompt_evidence(candidate, CASE_BY_ID, FIXTURE_DIGESTS)


def test_prompt_manifest_rejects_stale_section_hash(tmp_path: Path):
    steps = json.loads(STEPS_PATH.read_text(encoding="utf-8"))
    steps[5]["sha256"] = "0" * 64
    manifest = tmp_path / "system-steps.json"
    manifest.write_text(json.dumps(steps), encoding="utf-8")

    with pytest.raises(ValueError, match="changed. Update its hypothesis and evidence"):
        load_prompt_steps(manifest)


def test_prompt_assembly_excludes_provenance_but_preserves_instruction_comments(tmp_path: Path, monkeypatch):
    from nurse_scheduling.ai import system_prompt

    body = "Keep T literal.\n<!-- A model-facing instruction comment. -->"
    (tmp_path / "rule.md").write_text(
        "<!--\nSPDX-License-Identifier: AGPL-3.0-or-later\n-->\n"
        "<!-- This file is mostly AI generated. -->\n\n" + body + "\n"
    )
    monkeypatch.setattr(system_prompt, "PROMPT_DIRECTORY", tmp_path)
    assert system_prompt.load_system_prompt_sections([{"file": "rule.md"}]) == (body,)


def test_all_prompt_segments_have_headers_excluded_from_model_text():
    for step in load_prompt_steps():
        source = (PROMPT_DIRECTORY / step.file).read_text(encoding="utf-8")
        assert source.startswith("<!--\nThis file is part of Nurse Scheduling Project,")
        assert "SPDX-License-Identifier: AGPL-3.0-or-later\n-->\n<!-- This file is mostly AI generated. -->" in source
    assert "SPDX-License-Identifier:" not in SANDBOX_SYSTEM_PROMPT
    assert "This file is mostly AI generated." not in SANDBOX_SYSTEM_PROMPT
    assert "Copyright (C)" not in SANDBOX_SYSTEM_PROMPT


def test_print_prompt_script_matches_app_from_another_directory_without_dependencies(tmp_path: Path):
    script = PROMPT_DIRECTORY.parents[3] / "scripts/print_ai_system_prompt.py"
    result = subprocess.run(
        [sys.executable, "-S", str(script)], cwd=tmp_path, check=True, capture_output=True, text=True
    )
    assert result.stdout == SANDBOX_SYSTEM_PROMPT + "\n"
    assert not result.stderr


def test_prompt_comparison_defaults_to_three_repeats_and_step_cases():
    steps = load_prompt_steps()
    step_number = next(index for index, step in enumerate(steps, 1) if step.id == "resolve-ambiguous-targets")
    arguments, cases = _selected_cases(["--prompt-compare-step", str(step_number)])
    assert arguments.repeat == 3
    assert sorted(case.id for case in cases) == sorted(steps[step_number - 1].cases)


def test_optimizer_step_selects_its_direct_tool_case():
    step_number = next(index for index, step in enumerate(load_prompt_steps(), 1) if step.id == "optimizer-lifecycle")
    _, cases = _selected_cases(["--prompt-compare-step", str(step_number)])
    assert "tool-optimizer-start" in {case.id for case in cases}


@pytest.mark.parametrize(
    "argv",
    [
        ["--prompt-compare-step", "0"],
        ["--prompt-ablate-step", str(len(load_prompt_steps()) + 1)],
        ["--prompt-compare-step", "5", "--repeat", "2"],
        ["--prompt-compare-step", "5", "--repeat", "11"],
        ["--prompt-compare-step", "5", "--cost-ratio", "0.6"],
    ],
)
def test_prompt_comparison_rejects_invalid_requests(argv):
    with pytest.raises(SystemExit) as error:
        _parse_args(argv)
    assert error.value.code == 2


def test_prompt_comparison_allows_explicit_ten_run_investigation():
    arguments = _parse_args(["--prompt-compare-step", "1", "--repeat", "10"])
    assert arguments.repeat == 10


@pytest.mark.parametrize("argv", [[], ["--repeat", "3"], ["--jobs", "4"]])
def test_eval_cli_requires_explicit_scope_before_provider_work(argv, capsys):
    with pytest.raises(SystemExit) as error:
        main(argv)

    assert error.value.code == 2
    assert "choose --case, --category, or --tag" in capsys.readouterr().err


@pytest.mark.parametrize(
    "argv",
    [
        ["--full", "--case", "people-add"],
        ["--tuning", "--category", "03-structure"],
        ["--tuning", "--full"],
    ],
)
def test_eval_cli_rejects_broad_scope_combined_with_another_scope(argv):
    with pytest.raises(SystemExit) as error:
        _parse_args(argv)

    assert error.value.code == 2


@pytest.mark.parametrize("selector", ["--case", "--category", "--tag"])
def test_eval_cli_rejects_empty_selector(selector):
    with pytest.raises(SystemExit) as error:
        _parse_args([selector, ""])

    assert error.value.code == 2


def test_eval_cli_accepts_selected_tuning_and_full_scopes():
    assert _parse_args(["--case", "people-add", "--case", "people-group-members"]).case == [
        "people-add",
        "people-group-members",
    ]
    assert _parse_args(["--category", "01-reading"]).category == ["01-reading"]
    assert _parse_args(["--tag", "holdout"]).tag == ["holdout"]
    assert _parse_args(["--tuning"]).tuning
    assert _parse_args(["--full"]).full


def test_eval_cli_resolves_explicit_scopes_before_provider_work():
    _, selected = _selected_cases(["--case", "clarify-night-request-scope"])
    _, tuning = _selected_cases(["--tuning"])
    _, full = _selected_cases(["--full"])

    assert [case.id for case in selected] == ["clarify-night-request-scope"]
    assert 1 < len(tuning) < len(full)


@pytest.mark.parametrize(
    ("argv", "message"),
    [
        (["--case", "no-such-case"], "Unknown case ids"),
        (["--category", "no-such-category"], "Unknown categories"),
        (["--tag", "no-such-tag"], "Unknown tags"),
    ],
)
def test_eval_cli_rejects_unknown_scope_before_provider_work(argv, message):
    with pytest.raises(SystemExit, match=message):
        main(argv)


def settings(**overrides: object) -> AiSettings:
    values = {
        "provider_base_url": "https://provider.example/v1",
        "provider_api_key": "test-token",
        "provider_model": "test-model",
        "auth_token": "ai-shared-test-token",
    }
    values.update(overrides)
    return AiSettings(**values)


class ScriptedProvider:
    """Replay prepared turns, or fail, without contacting a provider."""

    def __init__(self, *turns) -> None:
        self._turns = list(turns)
        self.messages: list[Sequence[ChatMessage]] = []
        self.tool_definitions: list[Sequence[dict]] = []

    async def stream_events(self, messages: Sequence[ChatMessage], tools=None) -> AsyncIterator:
        self.messages.append(messages)
        self.tool_definitions.append(tools or ())
        turn = self._turns.pop(0) if self._turns else [TextDelta("Done.")]
        if isinstance(turn, Exception):
            raise turn
        for event in turn:
            yield event


class ConcurrentProvider:
    """Return one answer while measuring simultaneous provider streams."""

    def __init__(self) -> None:
        self.active = 0
        self.max_active = 0

    async def stream_events(self, messages: Sequence[ChatMessage], tools=None) -> AsyncIterator:
        self.active += 1
        self.max_active = max(self.max_active, self.active)
        try:
            await asyncio.sleep(0.01)
            yield TextDelta("Done.")
        finally:
            self.active -= 1


def _factory(command_handler=None) -> FakeSandboxFactory:
    if command_handler is None:
        command_handler = lambda *_: CommandResult("ok\n", "", 0)
    return FakeSandboxFactory(
        lambda sandbox_id: FakeSandboxBackend(sandbox_id, command_handler=command_handler),
    )


def _description_factory() -> FakeSandboxFactory:
    def edit(_command: str, _timeout: float | None, backend: FakeSandboxBackend) -> CommandResult:
        current = backend.files[WORKSPACE_SCHEDULE].decode()
        backend.files[WORKSPACE_SCHEDULE] = current.replace(
            "apiVersion: alpha\ndescription: ''",
            "apiVersion: alpha\ndescription: March ward roster",
        ).encode()
        return CommandResult("updated\n", "", 0)

    return _factory(edit)


def _run(case_id: str, provider: ScriptedProvider, factory: FakeSandboxFactory | None = None) -> CaseRun:
    return asyncio.run(run_case(provider, settings(), CASE_BY_ID[case_id], factory or _factory()))


def test_a_correct_answer_passes_and_records_its_cost():
    run = _run("ask-people-count", ScriptedProvider([TextDelta("There are 87 people.")]))

    assert run.passed
    assert run.category == "basics/00-summary"
    assert run.turns == 1
    assert run.tools == []
    assert not run.proposed
    assert run.seconds >= 0


def test_attachment_case_hydrates_generated_file_and_records_an_upload_event():
    factory = _factory()
    run = _run(
        "read-second-xlsx-sheet",
        ScriptedProvider(
            [ToolCallRequest((ToolCall("call-1", BASH_TOOL, '{"command":"inspect workbook"}'),))],
            [TextDelta("The code is NIGHT OWL 7429.")],
        ),
        factory,
    )

    assert run.passed
    backend = factory.created[0]
    assert backend.files["/workspace/attachments/01-ward-notes.xlsx"].startswith(b"PK")
    prompt = run.trajectory["prompt"]
    assert prompt[1]["content"].startswith("[App event] The user uploaded files.")
    assert '"path": "/workspace/attachments/01-ward-notes.xlsx"' in prompt[1]["content"]
    assert prompt[-1] == {"role": "user", "content": CASE_BY_ID["read-second-xlsx-sheet"].question}


def test_optimizer_case_uses_controlled_production_tool_contract():
    provider = ScriptedProvider(
        [ToolCallRequest((ToolCall("call-1", "optimizer", '{"action":"start"}'),))],
        [TextDelta("Started in the background. You can keep chatting while it runs.")],
    )
    run = _run("tool-optimizer-start", provider)

    assert run.passed
    assert run.tools == ["optimizer"]
    assert any(tool["function"]["name"] == "optimizer" for tool in provider.tool_definitions[0])
    assert optimizer_start_message(
        "eval-job", hashlib.sha256(fixture_text("small-clinic").encode()).hexdigest()
    ) == next(event["result"] for event in run.trajectory["events"] if event["kind"] == "tool")


def test_optimizer_status_control_requires_a_fresh_call_in_the_later_user_turn():
    case = CASE_BY_ID["tool-optimizer-status-on-request"]
    responses = [
        [ToolCallRequest((ToolCall("start", "optimizer", '{"action":"start"}'),))],
        [TextDelta("Optimization is running in the background. You can keep chatting.")],
    ]
    provider = ScriptedProvider(
        *responses,
        [ToolCallRequest((ToolCall("status", "optimizer", '{"action":"status"}'),))],
        [TextDelta("The optimizer is still running.")],
    )
    run = asyncio.run(run_case(provider, settings(), case, _factory()))
    assert run.passed
    assert run.tools == ["optimizer", "optimizer"]
    inferred = ScriptedProvider(*responses, [TextDelta("The optimizer is still running.")])
    run = asyncio.run(run_case(inferred, settings(), case, _factory()))
    assert not run.passed
    assert any("action" in failure and "status" in failure for failure in run.failures)


def test_optimizer_unavailability_is_a_tool_error_and_not_an_infrastructure_error():
    provider = ScriptedProvider(
        [ToolCallRequest((ToolCall("call-1", "optimizer", '{"action":"start"}'),))],
        [TextDelta("The optimizer API is unavailable.")],
    )
    run = _run("tool-optimizer-api-unavailable", provider)
    assert run.passed
    assert not run.error
    assert any(event.get("name") == "optimizer" and event.get("ok") is False for event in run.trajectory["events"])
    invented = _run("tool-optimizer-api-unavailable", ScriptedProvider([TextDelta("The API is unavailable.")]))
    assert not invented.passed
    assert "observes optimizer tool error" in "; ".join(invented.failures)


@pytest.mark.parametrize("fail_fast, expected_requests", [(True, 2), (False, 3)])
def test_optimizer_call_limit_stops_before_unnecessary_retries(fail_fast: bool, expected_requests: int):
    provider = ScriptedProvider(
        [ToolCallRequest((ToolCall("start", "optimizer", '{"action":"start"}'),))],
        [ToolCallRequest((ToolCall("status", "optimizer", '{"action":"status"}'),))],
        [TextDelta("The optimizer is unavailable.")],
    )
    run = asyncio.run(
        run_case(provider, settings(), CASE_BY_ID["tool-optimizer-api-unavailable"], _factory(), fail_fast=fail_fast)
    )
    assert not run.passed
    assert not run.error
    assert "uses optimizer at most 1 time(s): used 2" in run.failures
    assert len(provider.messages) == expected_requests
    assert any(event["kind"] == "evaluation_stop" for event in run.trajectory["events"]) == fail_fast


@pytest.mark.parametrize("fail_fast", [True, False])
def test_yaml_installation_violation_stops_before_executing_the_command(fail_fast):
    factory = _factory(lambda *_: CommandResult("", "Internet access is disabled.", 1))
    provider = ScriptedProvider(
        [ToolCallRequest((ToolCall("install", BASH_TOOL, '{"command":"pip install pyyaml"}'),))],
        [TextDelta("The generator could not run.")],
    )
    run = asyncio.run(
        run_case(
            provider,
            settings(),
            CASE_BY_ID["tool-yaml-generator-repair"],
            factory,
            fail_fast=fail_fast,
        )
    )
    assert not run.passed
    assert not run.error
    assert "avoids package installation" in " ".join(run.failures)
    assert len(factory.created[0].commands) == (0 if fail_fast else 1)
    assert factory.created[0].closed
    assert any(event["kind"] == "evaluation_stop" for event in run.trajectory["events"]) == fail_fast


def test_generator_installation_guard_does_not_change_other_case_trajectories():
    factory = _factory()
    _run(
        "tool-bash-count-weight",
        ScriptedProvider([ToolCallRequest((ToolCall("install", BASH_TOOL, '{"command":"pip install pyyaml"}'),))]),
        factory,
    )
    assert [command for command, _ in factory.created[0].commands] == ["pip install pyyaml"]


def test_missing_required_tool_does_not_stop_a_run_before_later_success():
    case = EvalCase(
        id="required-later",
        fixture="new-schedule",
        question="Inspect then start.",
        expect_proposal=False,
        tool_usage=ToolUsageExpectation(required=("optimizer",), max_total=2),
    )
    provider = ScriptedProvider(
        [ToolCallRequest((ToolCall("read", "read", '{"path":"/workspace/schedule.yaml"}'),))],
        [ToolCallRequest((ToolCall("start", "optimizer", '{"action":"start"}'),))],
        [TextDelta("Started in the background.")],
    )
    run = asyncio.run(run_case(provider, settings(), case, _factory()))
    assert run.passed
    assert len(provider.messages) == 3


def test_provider_wait_time_is_recorded_per_inference_turn():
    run = asyncio.run(run_case(ConcurrentProvider(), settings(), CASE_BY_ID["ask-people-count"], _factory()))

    assert run.llm_inference_seconds >= 0.005
    assert run.llm_turn_seconds == pytest.approx([run.llm_inference_seconds])


def test_provider_retries_are_reported_separately_from_logical_turns():
    run = _run(
        "ask-people-count",
        ScriptedProvider([ProviderAttempt(1), ProviderAttempt(2), TextDelta("There are 87 people.")]),
    )

    assert run.as_record()["provider_requests"] == {
        "turns": 1,
        "attempts": 2,
        "retries": 1,
        "retried_turns": 1,
        "attempts_per_turn": [2],
    }


def test_multi_tool_batches_are_recorded_per_model_turn():
    provider = ScriptedProvider(
        [
            ToolCallRequest(
                (
                    ToolCall("call_0", READ_TOOL, '{"path":"schedule.yaml"}'),
                    ToolCall("call_1", READ_TOOL, '{"path":"schedule.yaml"}'),
                )
            )
        ],
        [TextDelta("There are 87 people.")],
    )

    run = _run("ask-people-count", provider)

    assert run.tool_calls_per_turn == [2, 0]
    assert run.as_record()["tool_batches"] == {
        "count": 1,
        "multi_call_batches": 1,
        "max_calls_per_batch": 2,
        "calls_per_batch": [2],
        "calls_per_turn": [2, 0],
        "parallel_batches": 1,
        "parallel_per_batch": [True],
        "execution_seconds_per_batch": [pytest.approx(run.tool_batch_metrics[0].execution_seconds, abs=0.001)],
    }


def test_an_answer_in_words_is_accepted():
    run = _run("ask-people-count", ScriptedProvider([TextDelta("There are eighty-seven people.")]))

    assert run.passed


def test_a_wrong_answer_fails_with_the_reason():
    run = _run("ask-people-count", ScriptedProvider([TextDelta("There are 12 people.")]))

    assert not run.passed
    assert "not mentioned" in run.failures[0]


def test_an_edit_case_records_the_tools_and_the_proposal():
    provider = ScriptedProvider(
        [ToolCallRequest((ToolCall("call_0", BASH_TOOL, '{"command":"edit description"}'),))],
        [TextDelta("I propose the new description.")],
    )

    run = _run("description-set", provider, _description_factory())

    assert run.passed
    assert run.tools == [BASH_TOOL]
    assert run.proposed
    assert run.turns == 2


def test_run_case_passes_tool_activity_to_focused_trajectory_grading():
    case = EvalCase(
        id="focused-edit",
        fixture="new-schedule",
        question="Try an edit.",
        expect_proposal=False,
        tool_usage=ToolUsageExpectation(required=(EDIT_TOOL,)),
    )
    provider = ScriptedProvider(
        [ToolCallRequest((ToolCall("call_0", EDIT_TOOL, '{"path":"missing","edits":[]}'),))],
        [TextDelta("Done.")],
    )

    run = asyncio.run(run_case(provider, settings(), case, _factory()))

    assert not run.passed
    assert "not used successfully" in "; ".join(run.failures)


def test_eval_uses_the_sandbox_runner_and_closes_its_backend():
    factory = _description_factory()
    provider = ScriptedProvider(
        [ToolCallRequest((ToolCall("call_0", BASH_TOOL, '{"command":"edit"}'),))],
        [TextDelta("I propose the new description.")],
    )

    run = asyncio.run(
        run_case(
            provider,
            settings(),
            CASE_BY_ID["description-set"],
            factory,
        )
    )

    assert run.passed
    assert run.tools == [BASH_TOOL]
    assert run.proposed
    assert factory.created[0].closed
    assert "`/workspace/schedule.yaml`" in run.trajectory["prompt"][0]["content"]
    timing = run.as_record()["timing"]
    assert timing["end_to_end_seconds"] >= timing["llm_inference_seconds"]
    assert len(timing["llm_turn_seconds"]) == run.turns
    assert sum(timing["llm_turn_seconds"]) == pytest.approx(timing["llm_inference_seconds"], abs=0.002)
    assert timing["sandbox"]["available"] is True
    sandbox = timing["sandbox"]
    assert sandbox["lifetime_seconds"] == pytest.approx(
        sandbox["provisioning_seconds"]
        + sandbox["execution_seconds"]
        + sandbox["pause_transition_seconds"]
        + sandbox["warm_waiting_seconds"]
        + sandbox["suspended_seconds"]
        + sandbox["resume_wait_seconds"]
        + sandbox["teardown_seconds"],
        abs=0.005,
    )
    assert timing["end_to_end_seconds"] >= sandbox["lifetime_seconds"]
    assert timing["sandbox"]["suspension"] == {
        "pause_count": 0,
        "pause_cancel_count": 0,
        "resume_count": 0,
    }


def test_a_failed_tool_call_is_recorded_as_such():
    provider = ScriptedProvider(
        [ToolCallRequest((ToolCall("call_0", BASH_TOOL, '{"command":"false"}'),))],
        [TextDelta("I could not find that text.")],
    )
    factory = _factory(lambda *_: CommandResult("", "not found", 1))

    run = _run("description-set", provider, factory)

    assert not run.passed
    assert run.tools == [f"{BASH_TOOL}(failed)"]
    assert not run.proposed


def test_a_command_that_raises_is_recorded_before_the_sandbox_failure():
    provider = ScriptedProvider(
        [ToolCallRequest((ToolCall("call_0", BASH_TOOL, '{"command":"slow command"}'),))],
    )

    def fail(*_args):
        raise SandboxError("sandbox command failed")

    run = _run("description-set", provider, _factory(fail))

    assert not run.passed
    assert run.error == "sandbox command failed"
    assert run.trajectory["events"][-1] == {
        "kind": "tool_start",
        "name": BASH_TOOL,
        "arguments": '{"command":"slow command"}',
    }


def test_tool_calls_continue_until_the_model_finishes():
    provider = ScriptedProvider(
        *[[ToolCallRequest((ToolCall(f"call_{index}", BASH_TOOL, '{"command":"true"}'),))] for index in range(4)],
        [TextDelta("I need more input.")],
    )

    run = _run("ask-people-count", provider)

    assert run.tools == [BASH_TOOL] * 4
    assert run.turns == 5
    tool_events = [event for event in run.trajectory["events"] if event["kind"] == "tool"]
    assert len(tool_events) == 4


def test_a_provider_failure_is_reported_rather_than_raised():
    run = _run("ask-people-count", ScriptedProvider(ProviderError("The AI provider is unavailable.")))

    assert not run.passed
    assert run.error == "The AI provider is unavailable."
    assert run.failures == ["the provider failed"]


def test_reasoning_length_is_measured_without_entering_the_answer():
    provider = ScriptedProvider([ReasoningDelta("Counting people."), TextDelta("There are 87 people.")])

    run = _run("ask-people-count", provider)

    assert run.passed
    assert run.reasoning_chars == len("Counting people.")
    assert "Counting" not in run.answer


def test_token_usage_is_aggregated_across_provider_turns():
    provider = ScriptedProvider(
        [
            ToolCallRequest((ToolCall("call_0", BASH_TOOL, '{"command":"rg people"}'),)),
            TokenUsage(100, 20, 120, cached_prompt_tokens=40, reasoning_tokens=5),
        ],
        [TextDelta("There are 87 people."), TokenUsage(150, 10, 160, cached_prompt_tokens=90)],
    )

    run = _run("ask-people-count", provider)

    assert run.as_record()["token_usage"] == {
        "available": True,
        "complete": True,
        "reported_turns": 2,
        "prompt_tokens": 250,
        "cached_prompt_tokens": 130,
        "completion_tokens": 30,
        "reasoning_tokens": 5,
        "total_tokens": 280,
    }


def test_missing_provider_usage_is_recorded_explicitly():
    run = _run("ask-people-count", ScriptedProvider([TextDelta("There are 87 people.")]))

    assert run.as_record()["token_usage"] == {
        "available": False,
        "complete": False,
        "reported_turns": 0,
        "prompt_tokens": None,
        "cached_prompt_tokens": None,
        "completion_tokens": None,
        "reasoning_tokens": None,
        "total_tokens": None,
    }


def test_missing_cache_details_do_not_count_as_uncached_tokens():
    from tests.ai_eval.comparison import comparison_statistics
    from tests.ai_eval.runner import _prompt_cost

    run = _run("ask-people-count", ScriptedProvider([TextDelta("There are 87 people."), TokenUsage(100, 10, 110)]))

    assert run.as_record()["token_usage"]["cached_prompt_tokens"] is None
    assert _prompt_cost(run, "uncached-tokens") is None
    assert _prompt_cost(run, "total-tokens") == 110
    stats = comparison_statistics([run.as_record()], [run.as_record()])["cases"][run.case_id]["metrics"]
    assert stats["cached_prompt_tokens"]["pairs"] == 0
    assert stats["uncached_prompt_tokens"]["pairs"] == 0
    assert stats["total_tokens"]["pairs"] == 1


def test_cases_are_selected_by_id_and_by_category():
    cases = load_cases(CASES)

    assert {case.id for case in select(cases, [], [])} == {
        "answer-about-earlier-proposal",
        "apply-two-follow-up-edits",
        "approve-then-follow-up-edit",
        "cancel-ambiguous-request",
        "clarify-night-request-scope",
        "clarify-similar-people-groups",
        "confirm-one-cancel-other",
        "dates-range-expand-taiwan-detailed-yes",
        "dates-range-expand-taiwan-no",
        "dates-range-expand-taiwan-yes",
        "dates-range-shrink",
        "direct-combined-exact-edits",
        "download-generated-zip",
        "external-update-invalidates-pending",
        "explain-ai-attachments",
        "guide-add-person",
        "guide-load-yaml",
        "guide-run-optimization",
        "ask-people-group-union",
        "ask-positive-k-exceptions",
        "people-add-group-request",
        "people-group-members",
        "people-move-between-groups",
        "pref-add-two-student-requirements",
        "pref-copy-night-requests",
        "pref-copy-all-requests",
        "pref-copy-selected-shift-requests",
        "pref-group-affinity-specific-dates",
        "pref-modify-group-shift-count",
        "pref-modify-one-near-duplicate",
        "pref-requirement-modify-existing",
        "pronoun-selects-second-request",
        "question-then-approve-pending",
        "reject-contradicting-follow-up",
        "reject-conflicting-shift-request",
        "reject-shift-type-rename-collision",
        "reject-unknown-person",
        "reject-then-narrower-edit",
        "revise-unapproved-proposal-twice",
        "shift-type-remove-cascade",
        "shift-type-rename-cascade",
        "tool-write-minimal-schedule",
        "remember-edit-after-clarification",
        "reread-retained-upload",
        "revise-pending-copy-scope",
    }
    assert len(select(cases, [], [], full=True)) == len(cases)
    assert [case.id for case in select(cases, ["people-add"], [])] == ["people-add"]
    assert {case.category for case in select(cases, [], ["06-refusal"])} == {"basics/06-refusal"}
    assert {case.category for case in select(cases, [], ["basics/06-refusal"])} == {"basics/06-refusal"}
    assert {case.id for case in select(cases, [], [], ["taiwan-holidays"])} == {
        "dates-range-expand-taiwan-detailed-yes",
        "dates-range-expand-taiwan-no",
        "dates-range-expand-taiwan-yes",
        "heldout-expand-cross-year-no-renewal",
    }


def test_a_multi_user_turn_case_preserves_the_conversation_history():
    case = EvalCase(
        id="conversation",
        fixture="new-schedule",
        question="Expand the range.",
        expect_proposal=False,
        user_turns=("Expand the range.", "No."),
        intermediate_answer_contains=(("Taiwan",),),
    )
    provider = ScriptedProvider([TextDelta("Renew Taiwan holidays?")], [TextDelta("Okay, unchanged.")])

    run = asyncio.run(run_case(provider, settings(), case, _factory()))

    assert run.passed
    assert len(provider.messages) == 2
    assert provider.messages[1][-3:] == [
        {"role": "user", "content": "Expand the range."},
        {"role": "assistant", "content": "Renew Taiwan holidays?"},
        {"role": "user", "content": "No."},
    ]


def test_completion_only_case_seeds_history_without_counting_a_model_tool_call():
    case = CASE_BY_ID["result-single-tier-summary"]
    provider = ScriptedProvider([TextDelta('{"strongTotal":4,"strongUnmet":0}')])
    factory = _factory()
    run = asyncio.run(run_case(provider, settings(), case, factory))
    assert run.passed
    assert run.turns == 1
    assert run.tools == []
    assert len(provider.messages) == 1
    assert provider.messages[0][-4]["content"] == case.question
    assert "Optimization is running in the background" in provider.messages[0][-3]["content"]
    assert '"request_audit"' in provider.messages[0][-2]["content"]
    assert provider.messages[0][-1]["content"].startswith("[Current status]")
    assert not factory.created


@pytest.mark.parametrize("fail_fast, expected_user_turns", [(True, [1]), (False, [1, 2])])
def test_unexpected_proposal_fails_before_wasting_later_turns(fail_fast: bool, expected_user_turns: list[int]):
    provider = ScriptedProvider(
        [
            ToolCallRequest(
                (
                    ToolCall(
                        "edit-1",
                        "edit",
                        json.dumps(
                            {
                                "path": "/workspace/schedule.yaml",
                                "edits": [
                                    {
                                        "oldText": "description: Small pediatric clinic",
                                        "newText": "description: Premature",
                                    }
                                ],
                            }
                        ),
                    ),
                )
            )
        ],
        [TextDelta("Which shift type?")],
        [TextDelta("A later reply cannot undo the forbidden first proposal.")],
    )
    run = asyncio.run(
        run_case(
            provider, settings(), CASE_BY_ID["combined-edit-missing-shift-before-edit"], _factory(), fail_fast=fail_fast
        )
    )
    assert not run.passed
    assert not run.error
    assert "turn 1 proposal not expected" in "; ".join(run.failures)
    assert [event["turn"] for event in run.trajectory["events"] if event["kind"] == "user"] == expected_user_turns
    assert any(event["kind"] == "evaluation_stop" for event in run.trajectory["events"]) == fail_fast


def test_a_multi_user_turn_case_can_grade_an_earlier_proposal():
    case = EvalCase(
        id="conversation",
        fixture="new-schedule",
        question="Set the description.",
        expect_proposal=True,
        proposal_turn=1,
        user_turns=("Set the description.", "What is it now?"),
        expected_diff=(ExpectedDiff(path="description", before="", after="Ready", compares_value=True),),
        changes=("description",),
        answer_contains=("Ready",),
    )
    provider = ScriptedProvider(
        [ToolCallRequest((ToolCall("call_0", BASH_TOOL, '{"command":"edit description"}'),))],
        [TextDelta("I set it to Ready.")],
        [TextDelta("It is Ready.")],
    )

    def edit(_command: str, _timeout: float | None, backend: FakeSandboxBackend) -> CommandResult:
        current = backend.files[WORKSPACE_SCHEDULE].decode()
        backend.files[WORKSPACE_SCHEDULE] = current.replace("description: ''", "description: Ready", 1).encode()
        return CommandResult("updated\n", "", 0)

    run = asyncio.run(run_case(provider, settings(), case, _factory(edit)))

    assert run.passed


def test_approval_adopts_the_proposal_and_adds_trusted_history():
    case = EvalCase(
        id="approval",
        fixture="new-schedule",
        question="Set the description.",
        expect_proposal=True,
        proposal_turn=1,
        user_turns=("Set the description.", "Is it current?"),
        turn_actions=(TurnAction(1, "approve"),),
        expected_diff=(ExpectedDiff(path="description", before="", after="Ready", compares_value=True),),
        changes=("description",),
        answer_contains=("current",),
    )
    provider = ScriptedProvider(
        [ToolCallRequest((ToolCall("call_0", BASH_TOOL, '{"command":"edit description"}'),))],
        [TextDelta("I propose Ready.")],
        [TextDelta("It is current.")],
    )

    def edit(_command: str, _timeout: float | None, backend: FakeSandboxBackend) -> CommandResult:
        current = backend.files[WORKSPACE_SCHEDULE].decode()
        backend.files[WORKSPACE_SCHEDULE] = current.replace("description: ''", "description: Ready", 1).encode()
        return CommandResult("updated\n", "", 0)

    run = asyncio.run(run_case(provider, settings(), case, _factory(edit)))

    assert run.passed
    assert any("approved the previous schedule proposal" in message["content"] for message in provider.messages[2])


def test_an_unknown_case_id_stops_the_run():
    with pytest.raises(SystemExit, match="no-such-case"):
        select(load_cases(CASES), ["no-such-case"], [])


@pytest.mark.parametrize(("jobs", "expected_max_active"), [(1, 1), (2, 2)])
def test_run_all_bounds_parallelism_and_preserves_case_order(jobs: int, expected_max_active: int):
    cases = load_cases(CASES)[:3]
    provider = ConcurrentProvider()

    runs = asyncio.run(run_all(cases, settings(), provider, jobs, _factory()))

    assert provider.max_active == expected_max_active
    assert [run.case_id for run in runs] == [case.id for case in cases]


def test_run_all_rejects_non_positive_jobs():
    with pytest.raises(ValueError, match="jobs and repetitions must be positive"):
        asyncio.run(run_all([], settings(), ScriptedProvider(), 0, _factory()))


def test_run_all_repeats_cases_with_one_global_concurrency_limit():
    cases = load_cases(CASES)[:2]
    provider = ConcurrentProvider()

    runs = asyncio.run(run_all(cases, settings(), provider, 2, _factory(), repetitions=3))

    assert provider.max_active == 2
    assert [(run.case_id, run.repetition) for run in runs] == [
        (cases[0].id, 1),
        (cases[0].id, 2),
        (cases[0].id, 3),
        (cases[1].id, 1),
        (cases[1].id, 2),
        (cases[1].id, 3),
    ]


def test_completed_trace_is_saved_while_another_case_is_waiting_and_survives_cancellation(tmp_path, monkeypatch):
    from .ai_eval import runner

    cases = load_cases(CASES)[:2]
    checkpoints = runner.RunCheckpoints(tmp_path / "run", {"production": {"model": "test"}}, 1, 2)

    async def exercise():
        saved = asyncio.Event()
        waiting = asyncio.Event()
        cancelled = asyncio.Event()

        async def run_case(_provider, _settings, case, *args, **kwargs):
            if case.id == cases[0].id:
                return CaseRun(
                    case.id, case.category, False, 1, 1, error="provider unavailable", trajectory={"events": []}
                )
            waiting.set()
            try:
                await asyncio.Event().wait()
            finally:
                cancelled.set()

        def persist(run):
            checkpoints.record(run)
            saved.set()

        monkeypatch.setattr(runner, "run_case", run_case)
        task = asyncio.create_task(run_all(cases, settings(), ScriptedProvider(), 2, _factory(), on_complete=persist))
        await saved.wait()
        await waiting.wait()
        directory = tmp_path / "run"
        assert json.loads((directory / "cases" / f"{cases[0].id}.json").read_text())["error"] == "provider unavailable"
        assert not (directory / "cases" / f"{cases[1].id}.json").exists()
        assert json.loads((directory / "progress.json").read_text()) == {
            "status": "running",
            "completed": 1,
            "expected": 2,
        }
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
        assert cancelled.is_set()
        assert len((directory / "results.jsonl").read_text().splitlines()) == 1

    asyncio.run(exercise())


def test_checkpoints_separate_variants_and_repetitions_and_refuse_existing_output(tmp_path):
    from .ai_eval import runner

    root = tmp_path / "comparison"
    checkpoints = runner.RunCheckpoints(root, {"before": {"prompt": "A"}, "after": {"prompt": "B"}}, 3, 6)
    for label in ("after", "before"):
        run = CaseRun("case", "category", True, 1, 1, repetition=2, prompt_variant=label)
        checkpoints.record(run)
        assert (root / label / "cases" / "case--run-2.json").exists()
        assert json.loads((root / label / "results.jsonl").read_text())["prompt_variant"] == label
    with pytest.raises(FileExistsError):
        runner.RunCheckpoints(root, {"production": {}}, 1, 1)


def test_interrupted_cli_retains_completed_attempts_and_initial_metadata(tmp_path, monkeypatch):
    from .ai_eval import runner

    async def interrupted(*args, on_complete, **kwargs):
        on_complete(CaseRun("ask-people-count", "00-summary", False, 1, 1, error="provider unavailable"))
        raise RuntimeError("batch interrupted")

    monkeypatch.setattr(runner.AiSettings, "from_env", settings)
    monkeypatch.setattr(runner, "create_sandbox_factory", lambda _: FakeSandboxFactory())
    monkeypatch.setattr(runner, "run_all", interrupted)
    root = tmp_path / "run"
    with pytest.raises(RuntimeError, match="batch interrupted"):
        main(["--case", "ask-people-count", "--output-dir", str(root)])
    assert json.loads((root / "progress.json").read_text()) == {"status": "interrupted", "completed": 1, "expected": 1}
    assert json.loads((root / "metadata.json").read_text())["cases_sha256"]
    assert json.loads((root / "results.jsonl").read_text())["error"] == "provider unavailable"
    assert not (root / "summary.md").exists()


def test_the_summary_reports_each_category_and_every_failure():
    runs = [
        CaseRun("a", "00-summary", True, 2.0, 1, [], provider_attempts=2, provider_attempts_per_turn=[2]),
        CaseRun("b", "01-reading", False, 7.0, 2, [BASH_TOOL], ["answer mentions '27': not mentioned"]),
    ]

    report = summarize(runs)

    assert "00-summary         1/1" in report
    assert "01-reading         0/1" in report
    assert "total              1/2" in report
    assert "attempts" in report
    assert "retries" in report
    assert "b: answer mentions '27': not mentioned" in report
    assert summarize([]) == "No cases ran."


def test_the_report_records_enough_to_explain_a_run():
    run = _run("ask-people-count", ScriptedProvider([TextDelta("There are 87 people.")]))

    record = run.as_record()

    assert set(record) == {
        "case_id",
        "repetition",
        "prompt_variant",
        "category",
        "passed",
        "seconds",
        "timing",
        "turns",
        "tools",
        "failures",
        "proposed",
        "reasoning_chars",
        "answer",
        "error",
        "token_usage",
        "provider_requests",
        "tool_batches",
    }
    assert record["timing"]["end_to_end_seconds"] == pytest.approx(run.seconds, abs=0.001)
    assert record["timing"]["llm_inference_seconds"] == pytest.approx(run.llm_inference_seconds, abs=0.001)
    assert record["timing"]["llm_turn_seconds"] == pytest.approx(run.llm_turn_seconds, abs=0.001)
    assert record["timing"]["sandbox"]["available"] is True
    assert record["timing"]["sandbox"]["lifetime_seconds"] >= 0
    assert json.loads(json.dumps(record))["case_id"] == "ask-people-count"


def test_each_run_reports_into_its_own_timestamped_directory(monkeypatch: pytest.MonkeyPatch, tmp_path: Path):
    monkeypatch.setenv("AI_EVAL_ARTIFACT_ROOT", str(tmp_path))

    directory = default_output_dir()

    assert directory.parent == tmp_path / "ai-evals"
    assert directory.name.endswith("Z")


def test_a_report_holds_the_summary_and_one_line_for_each_case(tmp_path: Path):
    runs = [
        CaseRun("a", "00-summary", True, 2.0, 1, []),
        CaseRun("b", "01-reading", False, 7.0, 2, [BASH_TOOL], ["answer mentions '27': not mentioned"]),
    ]

    summary = write_report(runs, tmp_path / "run")

    assert summary.name == "summary.md"
    assert "00-summary         1/1" in summary.read_text(encoding="utf-8")
    lines = (tmp_path / "run" / "results.jsonl").read_text(encoding="utf-8").splitlines()
    assert [json.loads(line)["case_id"] for line in lines] == ["a", "b"]
    assert json.loads(lines[0])["token_usage"]["total_tokens"] is None


def test_a_report_records_case_concurrency_and_wall_time(tmp_path: Path):
    summary = write_report(
        [CaseRun("a", "00-summary", True, 2.0, 1, [])],
        tmp_path / "run",
        jobs=4,
        wall_seconds=1.25,
    )

    text = summary.read_text(encoding="utf-8")
    assert "Case concurrency: 4" in text
    assert "Wall time: 1.2 seconds" in text


def test_repeated_report_records_stability_and_distinct_trajectories(tmp_path: Path):
    runs = [
        CaseRun("a", "00-summary", True, 2.0, 2, [], repetition=1),
        CaseRun("a", "00-summary", False, 4.0, 4, [], error="provider failed", repetition=2),
    ]

    summary = write_report(runs, tmp_path / "run")

    text = summary.read_text(encoding="utf-8")
    assert "| a | 1/2 | 1 | 3.0 | 4.0 |" in text
    assert (tmp_path / "run/cases/a--run-1.json").exists()
    assert (tmp_path / "run/cases/a--run-2.json").exists()


def test_report_compares_reliability_and_cost_with_a_baseline(tmp_path: Path):
    baseline = tmp_path / "baseline"
    write_report([CaseRun("a", "00-summary", False, 4.0, 4, [])], baseline)
    current = CaseRun("a", "00-summary", True, 2.0, 2, [])

    summary = write_report([current], tmp_path / "current", baseline_report=baseline)

    assert "| a | 0% | 100% | +100% | 0 |" in summary.read_text(encoding="utf-8")
    statistics = json.loads((tmp_path / "current/baseline-comparison.json").read_text())
    assert statistics["cases"]["a"]["passing_pairs"] == 0
    assert statistics["cases"]["a"]["metrics"]["total_tokens"]["delta"]["mean"] is None


def test_prompt_comparison_requires_after_passes_and_a_measured_gain():
    before = [
        CaseRun("a", "test", False, 2.0, 2, [], repetition=repetition, prompt_variant="before")
        for repetition in range(1, 4)
    ]
    after = [
        CaseRun("a", "test", True, 2.0, 2, [], repetition=repetition, prompt_variant="after")
        for repetition in range(1, 4)
    ]
    report, improved = prompt_comparison_markdown([*before, *after])
    assert improved
    assert "| a | 0/3 | 3/3 | 0 | n/a |" in report

    after[0].passed = False
    assert not prompt_comparison_markdown([*before, *after])[1]
    after[0].passed = True
    after[0].error = "sandbox failed"
    assert not prompt_comparison_markdown([*before, *after])[1]


def test_prompt_comparison_cost_gain_requires_explicit_ratio():
    before = [
        CaseRun("a", "test", True, 2.0, 3, [READ_TOOL] * 5, repetition=repetition, prompt_variant="before")
        for repetition in range(1, 4)
    ]
    after = [
        CaseRun("a", "test", True, 2.0, 3, [READ_TOOL] * 2, repetition=repetition, prompt_variant="after")
        for repetition in range(1, 4)
    ]
    runs = [*before, *after]
    assert not prompt_comparison_markdown(runs)[1]
    report, improved = prompt_comparison_markdown(runs, "tool-calls", 0.6)
    assert improved
    assert "| a | 3/3 | 3/3 | 0 | 0.40 |" in report

    other_before = [
        CaseRun("b", "test", True, 2.0, 3, [READ_TOOL] * 2, repetition=repetition, prompt_variant="before")
        for repetition in range(1, 4)
    ]
    other_after = [
        CaseRun("b", "test", True, 2.0, 3, [READ_TOOL] * 4, repetition=repetition, prompt_variant="after")
        for repetition in range(1, 4)
    ]
    assert not prompt_comparison_markdown([*runs, *other_before, *other_after], "tool-calls", 0.6)[1]


def test_prompt_comparison_rejects_cost_regression_in_another_case():
    improved_before = [
        CaseRun("a", "test", False, 2.0, 2, [], repetition=repetition, prompt_variant="before")
        for repetition in range(1, 4)
    ]
    improved_after = [
        CaseRun("a", "test", True, 2.0, 2, [], repetition=repetition, prompt_variant="after")
        for repetition in range(1, 4)
    ]
    costly_before = [
        CaseRun("b", "test", True, 2.0, 2, [READ_TOOL], repetition=repetition, prompt_variant="before")
        for repetition in range(1, 4)
    ]
    costly_after = [
        CaseRun("b", "test", True, 2.0, 2, [READ_TOOL] * 2, repetition=repetition, prompt_variant="after")
        for repetition in range(1, 4)
    ]

    report, improved = prompt_comparison_markdown(
        [*improved_before, *improved_after, *costly_before, *costly_after], "tool-calls", 0.8
    )

    assert not improved
    assert "Decision: cost regression observed." in report


def _token_comparison_runs():
    runs = []
    for variant, completions in (("before", (100, 200, 300)), ("after", (90, 150, 120))):
        for repetition, completion in enumerate(completions, 1):
            runs.append(
                CaseRun(
                    "a",
                    "test",
                    True,
                    2.0,
                    2,
                    [READ_TOOL, f"{BASH_TOOL}(failed)"],
                    repetition=repetition,
                    prompt_variant=variant,
                    token_usage=TokenUsage(1000, completion, 1000 + completion, 200, completion // 2),
                    token_usage_turns=2,
                )
            )
    return runs


def test_comparison_reports_paired_token_means_and_sample_standard_deviation():
    runs = _token_comparison_runs()
    # Arrival order differs from repetition order in concurrent evaluations.
    statistics = comparison_statistics(
        [run.as_record() for run in runs[:3]],
        [run.as_record() for run in reversed(runs[3:])],
    )
    case = statistics["cases"]["a"]
    completion = case["metrics"]["completion_tokens"]
    assert completion["before"] == {"mean": 200.0, "std": 100.0}
    assert completion["after"] == {"mean": 120.0, "std": 30.0}
    assert completion["delta"]["mean"] == -80.0
    assert completion["delta"]["std"] == pytest.approx(7900**0.5)
    assert completion["relative_delta"] == pytest.approx(-0.4)
    assert case["metrics"]["uncached_prompt_tokens"]["before"]["mean"] == 800
    assert case["metrics"]["bash_calls"]["before"]["mean"] == 1
    report, improved = prompt_comparison_markdown(runs)
    assert not improved  # Descriptive gains do not silently select a benefit target.
    assert "| a | completion_tokens | 3/3 | 200.0 ± 100.0 | 120.0 ± 30.0 | -80.0 ± 88.9 | -40.0% |" in report
    assert "| a | total_tokens | 3/3 | 1200.0 ± 100.0 | 1120.0 ± 30.0 | -80.0 ± 88.9 |" in report


@pytest.mark.parametrize("usage_problem", ["partial", "missing"])
def test_comparison_excludes_failed_pairs_and_incomplete_usage_without_losing_reliability(usage_problem):
    runs = _token_comparison_runs()
    runs[0].passed = False
    runs[4].passed = False
    runs[4].error = "sandbox failed"
    runs[5].token_usage_turns = 1
    if usage_problem == "missing":
        runs[5].token_usage = None
    statistics = comparison_statistics([run.as_record() for run in runs[:3]], [run.as_record() for run in runs[3:]])
    case = statistics["cases"]["a"]
    assert case["before"]["passed"] == case["after"]["passed"] == 2
    assert case["after"]["infrastructure_errors"] == 1
    assert case["passing_pairs"] == 1
    assert case["metrics"]["total_tokens"]["pairs"] == 0
    assert case["metrics"]["total_tokens"]["missing_pairs"] == 1
    assert case["metrics"]["total_tokens"]["delta"] == {"mean": None, "std": None}
    assert case["metrics"]["reads"]["pairs"] == 1
    assert "| a | total_tokens | 0/3 | n/a | n/a | n/a | n/a |" in comparison_metrics_markdown(statistics)


def test_comparison_zero_tokens_single_pair_and_unmatched_runs_are_explicit():
    run = CaseRun("a", "test", True, 2.0, 1, [], token_usage=TokenUsage(0, 0, 0), token_usage_turns=1)
    before = [run.as_record(), replace(run, repetition=2).as_record()]
    statistics = comparison_statistics(before, [run.as_record()])
    case = statistics["cases"]["a"]
    assert case["before"]["runs"] == 2
    assert case["after"]["runs"] == 1
    assert case["unmatched_before"] == 1
    tokens = case["metrics"]["total_tokens"]
    assert tokens["before"] == tokens["delta"] == {"mean": 0, "std": None}
    assert tokens["relative_delta"] is None
    report = comparison_metrics_markdown(statistics)
    assert "| a | total_tokens | 1/1 | 0.0 (SD n/a) | 0.0 (SD n/a) | +0.0 (SD n/a) | n/a |" in report
    assert "unmatched repetitions excluded: before 1, after 0" in report


def test_comparison_rejects_duplicate_and_mismatched_repetitions():
    runs = _token_comparison_runs()
    record = runs[0].as_record()
    with pytest.raises(ValueError, match="Duplicate case/repetition"):
        comparison_statistics([record, record], [record])
    runs[-1].repetition = 4
    with pytest.raises(ValueError, match="Unpaired prompt repetitions"):
        prompt_comparison_markdown(runs)


@pytest.mark.parametrize("metric", ["completion-tokens", "total-tokens"])
def test_prompt_comparison_token_targets_require_complete_usage(metric):
    runs = _token_comparison_runs()
    target = 0.7 if metric == "completion-tokens" else 0.95
    _parse_args(["--prompt-compare-step", "1", "--cost-metric", metric, "--cost-ratio", str(target)])
    assert prompt_comparison_markdown(runs, metric, target)[1]
    runs[-1].token_usage_turns = 1
    report, improved = prompt_comparison_markdown(runs, metric, target)
    assert not improved
    assert "missing cost data" in report


def test_prompt_comparison_cli_writes_first_class_metrics_without_a_cost_target(tmp_path: Path, monkeypatch):
    from .ai_eval import runner

    async def recorded_runs(*args, **kwargs):
        return [replace(run, case_id="ask-people-count") for run in _token_comparison_runs()]

    monkeypatch.setattr(runner.AiSettings, "from_env", settings)
    monkeypatch.setattr(runner, "create_sandbox_factory", lambda _: FakeSandboxFactory())
    monkeypatch.setattr(runner, "run_all", recorded_runs)
    output = tmp_path / "comparison"
    exit_code = main(["--prompt-compare-step", "1", "--case", "ask-people-count", "--output-dir", str(output)])
    assert exit_code == 1
    statistics = json.loads((output / "comparison.json").read_text())
    assert not statistics["improved"]
    assert statistics["cost_metric"] is None
    case = statistics["cases"]["ask-people-count"]
    assert case["before"]["pass_rate"] == case["after"]["pass_rate"] == 1
    assert case["metrics"]["completion_tokens"]["delta"]["mean"] == -80
    assert "-80.0 ± 88.9" in (output / "comparison.md").read_text(encoding="utf-8")


def test_run_all_injects_each_prompt_variant_into_every_repetition():
    provider = ScriptedProvider(*[[TextDelta("There are 87 people.")] for _ in range(6)])
    variants = (("before", prompt_at_step(0)), ("after", prompt_at_step(1)))
    runs = asyncio.run(
        run_all(
            [CASE_BY_ID["ask-people-count"]],
            settings(),
            provider,
            1,
            _factory(),
            3,
            prompt_variants=variants,
        )
    )
    assert all(run.passed for run in runs)
    assert [(run.repetition, run.prompt_variant) for run in runs] == [
        (1, "before"),
        (1, "after"),
        (2, "after"),
        (2, "before"),
        (3, "before"),
        (3, "after"),
    ]
    for run in runs:
        actual = run.trajectory["prompt"][0]["content"]
        assert actual.startswith(dict(variants)[run.prompt_variant])
        assert "Current schedule summary:" in actual


def test_reference_digests_cover_every_file_hydrated_into_the_sandbox():
    digests = _reference_digests()

    # A user guide edit steers the 10-app-ui cases, so it has to move this fingerprint.
    expected = {path.name for path in SCHEMA_REFERENCE_FILES.values()}
    expected.add(TAIWAN_HOLIDAYS_SOURCE.name)
    expected.update(f"user-guide/{relative}" for relative in load_user_guide_references())
    expected.update(path.removeprefix("/reference/") for path in REFERENCE_ATTACHMENT_TOOLS)
    expected.update({"tools/README.md", "result_context.py", "policy_audit"})
    assert set(digests) == expected
    assert digests == dict(sorted(digests.items()))
    assert all(len(digest) == 64 for digest in digests.values())


def test_report_writes_reproducibility_metadata(tmp_path: Path):
    metadata = {"git_revision": "abc", "prompt_sha256": "123"}

    write_report([CaseRun("a", "00-summary", True, 2.0, 1, [])], tmp_path / "run", metadata=metadata)

    assert json.loads((tmp_path / "run/metadata.json").read_text(encoding="utf-8")) == metadata


def test_metadata_hashes_selected_prompt_and_case_criteria():
    case = CASE_BY_ID["ask-people-count"]
    before = _evaluation_metadata(settings(), [case], 3, prompt_at_step(0))
    after = _evaluation_metadata(settings(), [case], 3, prompt_at_step(1))
    changed_case = _evaluation_metadata(
        settings(), [replace(case, question="A different question")], 3, prompt_at_step(0)
    )

    assert before["prompt_sha256"] != after["prompt_sha256"]
    assert before["cases_sha256"][case.id] == after["cases_sha256"][case.id]
    assert before["cases_sha256"][case.id] != changed_case["cases_sha256"][case.id]


def test_summary_markdown_reports_every_sandbox_metric_per_case(tmp_path: Path):
    metrics = SandboxTurnMetrics(
        provisioning_seconds=0.4,
        execution_seconds=1.0,
        pause_transition_seconds=0.3,
        warm_waiting_seconds=3.0,
        suspended_seconds=5.0,
        resume_wait_seconds=0.1,
        max_resume_wait_seconds=0.06,
        teardown_seconds=0.2,
        lifetime_seconds=10.0,
        pause_count=2,
        pause_cancel_count=1,
        resume_count=2,
    )
    summary = write_report(
        [CaseRun("a", "00-summary", True, 10.0, 1, [], sandbox_metrics=metrics)],
        tmp_path / "run",
    )

    text = summary.read_text(encoding="utf-8")
    assert "mutually exclusive lifetime components" in text
    assert "| a | 10.000 | 0.400 | 1.000 | 0.300 | 3.000 | 5.000 | 0.100 | 0.200 |" in text
    assert "| a | 2 | 1 | 2 | 0.100 | 0.060 |" in text
    assert "| a | 0 | 0 | 0 | none | 0 | none |" in text


def test_a_report_never_overwrites_an_earlier_one(tmp_path: Path):
    write_report([CaseRun("a", "00-summary", True, 2.0, 1, [])], tmp_path / "run")

    with pytest.raises(FileExistsError):
        write_report([CaseRun("a", "00-summary", True, 2.0, 1, [])], tmp_path / "run")


def test_a_run_records_everything_it_did():
    edit = '{"command":"edit description"}'
    provider = ScriptedProvider(
        [
            ReasoningDelta("Looking for the description. "),
            ToolCallRequest((ToolCall("call_0", BASH_TOOL, edit),)),
        ],
        [TextDelta("I propose the new description.")],
    )

    trajectory = _run("description-set", provider, _description_factory()).as_trajectory()

    assert trajectory["question"].startswith("Give this schedule the description")
    assert trajectory["prompt"][0]["role"] == "system"
    kinds = [event["kind"] for event in trajectory["events"]]
    assert kinds == ["user", "reasoning", "tool_start", "tool", "text", "proposal"]
    assert trajectory["events"][2]["arguments"] == edit
    tool_event = trajectory["events"][3]
    assert tool_event["ok"]
    assert "passed trusted server-side validation" in tool_event["result"]
    assert "March ward roster" in trajectory["proposal"]["schedule_yaml"]
    assert all(check["passed"] for check in trajectory["checks"])


def test_a_report_keeps_one_trajectory_file_for_each_case(tmp_path: Path):
    runs = [CaseRun("a", "00-summary", True, 2.0, 1, [], trajectory={"events": [{"kind": "text", "text": "hi"}]})]

    write_report(runs, tmp_path / "run")

    trajectory = json.loads((tmp_path / "run" / "cases" / "a.json").read_text(encoding="utf-8"))
    assert trajectory["case_id"] == "a"
    assert trajectory["events"] == [{"kind": "text", "text": "hi"}]


def test_command_timeout_keeps_original_tool_failure_and_is_not_an_infrastructure_error():
    def timeout(_command, _timeout, backend):
        backend.closed = True
        return CommandResult("", "", 124, timed_out=True, sandbox_terminated=True)

    provider = ScriptedProvider(
        [ToolCallRequest((ToolCall("wait", BASH_TOOL, '{"command":"sleep 30"}'),))],
    )
    run = asyncio.run(run_case(provider, settings(), CASE_BY_ID["tool-optimizer-start"], _factory(timeout)))
    assert not run.passed
    assert not run.error
    assert run.tools == [f"{BASH_TOOL}(failed)"]
    assert "Command timed out" in run.failures[0]
    outcome = next(event for event in run.trajectory["events"] if event["kind"] == "tool")
    assert not outcome["ok"]
    assert "Command timed out" in outcome["result"]
    assert any(event["kind"] == "evaluation_stop" for event in run.trajectory["events"])
    assert len(provider.messages) == 1


def test_recoverable_command_timeout_remains_in_a_successful_evaluation_trajectory():
    provider = ScriptedProvider(
        [ToolCallRequest((ToolCall("wait", BASH_TOOL, '{"command":"sleep 30"}'),))],
        [ToolCallRequest((ToolCall("start", "optimizer", '{"action":"start"}'),))],
        [TextDelta("Started in the background. You can keep chatting while it runs.")],
    )
    run = asyncio.run(
        run_case(
            provider,
            settings(),
            CASE_BY_ID["tool-optimizer-start"],
            _factory(lambda *_: CommandResult("partial\n", "", 124, timed_out=True)),
        )
    )
    assert run.passed
    assert not run.error
    assert run.tools == [f"{BASH_TOOL}(failed)", "optimizer"]
    assert "Command timed out" in next(event["result"] for event in run.trajectory["events"] if event["kind"] == "tool")


@pytest.mark.parametrize(
    "case_id,relevant",
    [
        ("pdf-visual-layout", "inspect_pdf.py"),
        ("xlsx-styled-requests-and-caches", "inspect_xlsx.py"),
        ("result-assignment-details", "inspect_optimizer_result.py"),
    ],
)
def test_helper_receipts_ignore_unrelated_scripts_but_bind_relevant_content(case_id, relevant, tmp_path, monkeypatch):
    from nurse_scheduling.ai.sandbox_agent import INSPECTION_HELPERS

    case = CASE_BY_ID[case_id]
    before = case_digest(case)
    unrelated = tmp_path / "unrelated.py"
    unrelated.write_text("# Independent helper\n")
    monkeypatch.setitem(REFERENCE_ATTACHMENT_TOOLS, "/reference/tools/unrelated.py", unrelated)
    monkeypatch.setitem(INSPECTION_HELPERS, "unrelated.py", "Independent capability")
    assert case_digest(case) == before
    changed = tmp_path / relevant
    changed.write_text("# Changed relevant helper\n")
    monkeypatch.setitem(REFERENCE_ATTACHMENT_TOOLS, f"/reference/tools/{relevant}", changed)
    assert case_digest(case) != before


def test_request_receipt_binds_compiled_selector_producer(monkeypatch):
    from nurse_scheduling.ai import result_context

    case = CASE_BY_ID["request-tier-counts-large"]
    before = case_digest(case)
    pdf_before = case_digest(CASE_BY_ID["pdf-visual-layout"])

    def replacement(data, source):
        return {}

    monkeypatch.setattr(result_context, "_project_context", replacement)
    assert case_digest(case) != before
    assert case_digest(CASE_BY_ID["pdf-visual-layout"]) == pdf_before


@pytest.mark.parametrize("fixture", sorted(set(FIXTURES) - {"new-schedule"}))
def test_complete_evaluation_fixtures_match_frontend_subset(fixture):
    # new-schedule intentionally starts with an incomplete construction scaffold.
    result = validate_frontend_schedule_yaml(fixture_text(fixture), max_bytes=2000000)
    assert result.valid, result.issues


@pytest.mark.parametrize("response_limit", [True, False], ids=["model-limit", "provider-outage"])
def test_response_limit_is_a_behavior_failure_and_outages_remain_infrastructure(response_limit):
    error = (
        ProviderResponseLimitError.for_user("The AI provider returned more reasoning than one answer may contain.")
        if response_limit
        else ProviderError("connection unavailable")
    )

    class FailingProvider:
        async def stream_events(self, messages, tools=None):
            yield ReasoningDelta("Checking the source.")
            raise error

    run = asyncio.run(run_case(FailingProvider(), settings(), CASE_BY_ID["ask-people-count"], _factory()))
    assert not run.passed
    assert bool(run.error) is not response_limit
    assert run.reasoning_chars == len("Checking the source.")
    if response_limit:
        assert run.failures == [str(error)]
        assert run.trajectory["events"][-1] == {"kind": "evaluation_stop", "reason": str(error)}
    else:
        assert run.failures == ["the provider failed"]


def test_explained_workbook_case_accepts_valid_import_and_rejects_changed_weights():
    source = """apiVersion: alpha
description: November 2025
dates:
  range:
    startDate: '2025-11-01'
    endDate: '2025-11-03'
  items: []
  groups: []
people:
  items:
  - id: Ada
    description: HN
    history:
    - OFF
  - id: Bela
    description: N
    history:
    - E
    - E
  groups: []
shiftTypes:
  items:
  - id: D
    description: Day
  - id: E
    description: Evening
  groups: []
preferences:
- type: at most one shift per day
- type: shift request
  person:
  - Ada
  date:
  - '01'
  shiftType:
  - OFF
  weight: 11000000000
- type: shift request
  person:
  - Ada
  date:
  - '03'
  shiftType:
  - D
  weight: 11000000000
- type: shift request
  person:
  - Bela
  date:
  - '01'
  shiftType:
  - OFF
  weight: 11000000
- type: shift request
  person:
  - Bela
  date:
  - '02'
  shiftType:
  - E
  weight: 11000000000
"""
    assert validate_frontend_schedule_yaml(source, max_bytes=2_000_000).valid
    case = CASE_BY_ID["workbook-explained-policy"]
    initial = _load_yaml(fixture_text(case.fixture).encode())
    proposal = YAML(typ="safe").load(source)
    correct = grade(case, RunOutcome(answer="", initial=initial, proposed=proposal, activity=[]))
    assert correct.passed, correct.failures()
    changed = deepcopy(proposal)
    changed["preferences"][3]["weight"] = 11_000_000_000
    wrong = grade(case, RunOutcome(answer="", initial=initial, proposed=changed, activity=[]))
    assert not wrong.passed
    changed["preferences"][3]["weight"] = float("inf")
    mandatory = grade(case, RunOutcome(answer="", initial=initial, proposed=changed, activity=[]))
    assert not mandatory.passed


def test_history_import_case_checks_every_source_history():
    from io import BytesIO

    from openpyxl import load_workbook

    from .ai_eval.attachment_fixtures import load_attachment_fixtures

    case = CASE_BY_ID["workbook-preserve-previous-history"]
    book = load_workbook(BytesIO(load_attachment_fixtures(case.attachments)[0].data))
    people = []
    for row in book.active.iter_rows(min_row=4, max_col=3, values_only=True):
        _, code, person = row
        if not person:
            continue
        if code.lower() == "off":
            history = ["OFF"]
        else:
            history = [code[-1]] * int(code[:-1] or 1)
        people.append({"id": person, "description": "", "history": history})
    book.close()
    initial = _load_yaml(fixture_text(case.fixture).encode())
    proposal = deepcopy(initial)
    proposal["people"]["items"] = people
    correct = grade(case, RunOutcome(answer="", initial=initial, proposed=proposal, activity=[]))
    assert correct.passed, correct.failures()
    for person in proposal["people"]["items"]:
        if person["history"] == ["OFF"]:
            person["history"] = []
    wrong = grade(case, RunOutcome(answer="", initial=initial, proposed=proposal, activity=[]))
    assert len(wrong.failures()) == 31
