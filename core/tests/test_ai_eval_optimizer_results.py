"""Controlled optimizer result delivery and semantic grading checks."""

# This file is part of Nurse Scheduling Project, see <https://github.com/j3soon/nurse-scheduling>.
#
# Copyright (C) 2023-2026 Johnson Sun
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as
# published by the Free Software Foundation, either version 3 of the
# License, or (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.

# This test is mostly AI generated.

import asyncio
import hashlib
import json
from dataclasses import replace
from io import BytesIO

import pytest
from openpyxl import load_workbook

from nurse_scheduling.ai.optimizer import WORKSPACE_OPTIMIZER_RESULT, optimizer_completion_message
from nurse_scheduling.ai.optimizer_tool import OPTIMIZER_RUNNING_MESSAGE
from nurse_scheduling.ai.provider import TextDelta, ToolCall, ToolCallRequest
from nurse_scheduling.ai.validation import validate_frontend_schedule_yaml
from nurse_scheduling.ai.workspace import WORKSPACE_RESULT_CONTEXT, WORKSPACE_SCHEDULE

from .ai_eval.grading import EvalCase, RunOutcome, grade
from .ai_eval.optimizer_fixtures import ASSIGNMENTS, FIXTURE, completion_result
from .ai_eval.prompt_ladder import case_digest, load_prompt_steps, validate_prompt_evidence
from .ai_eval.runner import run_case
from .test_ai_eval_runner import CASE_BY_ID, FIXTURE_DIGESTS, ScriptedProvider, _factory, settings


def test_controlled_completion_exports_the_verified_assignment():
    source = FIXTURE.read_text(encoding="utf-8")
    assert validate_frontend_schedule_yaml(source, 1_000_000).valid
    workbook, metadata = completion_result("request-audit", source)
    assert metadata["source_sha256"] == hashlib.sha256(source.encode()).hexdigest()
    assert metadata["result"]["score"] == 33_000_000_000
    ws = load_workbook(BytesIO(workbook)).active
    for row in range(3, 6):
        person = ws.cell(row, 1).value
        for col, assigned in enumerate(ASSIGNMENTS[person], 5):
            value = ws.cell(row, col).value or ""
            assignment = value.split(" [", 1)[0].strip() or "OFF"
            assert assignment == assigned
    assert ws.cell(7, 5).value == "FEASIBLE"
    notes = list(load_workbook(BytesIO(workbook))["Notes"].values)
    assert len(notes) == 3
    assert {r[2] for r in notes[1:]} == {
        "Weight of unmet single-style request: 11000000000",
        "Weight of unmet single-style request: 11000000",
    }


def test_completion_uses_the_production_message_and_mounts_a_fresh_result():
    case = CASE_BY_ID["optimizer-score-same-model"]
    provider = ScriptedProvider(
        [ToolCallRequest((ToolCall("start", "optimizer", '{"action":"start","timeout_seconds":60}'),))],
        [ToolCallRequest((ToolCall("inspect", "bash", '{"command":"inspect result"}'),))],
        [TextDelta(json.dumps(case.answer_json))],
    )
    factory = _factory()
    run = asyncio.run(run_case(provider, settings(), case, factory))
    assert run.passed, run.failures
    source = factory.created[0].files[WORKSPACE_SCHEDULE].decode()
    workbook, metadata = completion_result(case.optimizer_completion, source)
    assert WORKSPACE_OPTIMIZER_RESULT not in factory.created[0].files
    assert factory.created[1].files[WORKSPACE_OPTIMIZER_RESULT] == workbook
    completion = run.trajectory["prompts"][1][-2]
    assert completion == {"role": "user", "content": optimizer_completion_message(metadata)}
    assert run.trajectory["prompts"][1][-1]["content"].startswith("[Current status]")
    assert run.trajectory["prompts"][1][-3]["content"] == OPTIMIZER_RUNNING_MESSAGE
    assert all(backend.closed for backend in factory.created)
    assert any(e["kind"] == "optimizer" for e in run.trajectory["events"])


def test_completion_is_not_delivered_without_a_started_job():
    case = CASE_BY_ID["optimizer-score-same-model"]
    provider = ScriptedProvider([TextDelta(json.dumps(case.answer_json))])
    run = asyncio.run(run_case(provider, settings(), case, _factory()))
    assert not run.passed
    assert len(run.trajectory["prompts"]) == 1
    assert run.trajectory["events"][-1]["reason"] == "optimizer was not started"


def test_real_optimizer_solves_without_replaying_a_fixed_assignment(tmp_path):
    from .ai_eval.real_optimizer import solve
    from .ai_eval.runner import fixture_text

    source = fixture_text("policy-audit")
    receipt = solve(source, tmp_path, 5)
    metadata = receipt["metadata"]
    assert metadata["result"] == {
        "outcome": "optimal",
        "score": 44_011_000_000,
        "solver_status": "OPTIMAL",
        "termination_reason": "optimality_proven",
    }
    assert all(row["unmet"] == 0 for row in metadata["request_audit"]["summary"])
    assert metadata["request_audit"]["policy"]["staffing"][0]["preferred_shortfall"] == 0
    assert metadata["request_audit"]["policy"]["successions"][0]["unmet"] == 0
    assert (tmp_path / "submitted.yaml").read_text(encoding="utf-8") == source
    assert load_workbook(tmp_path / "result.xlsx").active.cell(8, 5).value == "OPTIMAL"


@pytest.mark.parametrize("conflict", [False, True])
def test_request_explanation_fixtures_distinguish_a_miss_from_a_proven_conflict(tmp_path, conflict):
    from .ai_eval.optimizer_fixtures import RESULT_SOURCES
    from .ai_eval.real_optimizer import solve

    name = "request-audit-conflict" if conflict else "request-audit-all-strong"
    source = RESULT_SOURCES[name].read_text(encoding="utf-8")
    assert validate_frontend_schedule_yaml(source, 1_000_000).valid
    _, incumbent = completion_result(name, source)
    weak = next(row for row in incumbent["request_audit"]["summary"] if row["weight"] == 11_000_000)
    assert weak["unmet"] == 1
    solved = solve(source, tmp_path, 5)["metadata"]
    assert solved["result"]["solver_status"] == "OPTIMAL"
    assert solved["result"]["score"] == (44_000_000_000 if conflict else 44_011_000_000)
    weak = next(row for row in solved["request_audit"]["summary"] if row["weight"] == 11_000_000)
    assert weak["unmet"] == int(conflict)


def test_real_completion_factory_receives_the_actual_submitted_source(tmp_path):
    from .ai_eval.real_optimizer import solve

    case = replace(
        CASE_BY_ID["optimizer-score-same-model"],
        fixture="policy-audit",
        answer_json={"score": 44_011_000_000},
        answer_contains=(),
        answer_matches=(),
        answer_not_matches=(),
    )
    provider = ScriptedProvider(
        [ToolCallRequest((ToolCall("start", "optimizer", '{"action":"start","timeout_seconds":60}'),))],
        [TextDelta('{"score":44011000000}')],
    )
    received = []

    def completion(name, source):
        received.append((name, source))
        receipt = solve(source, tmp_path, 5)
        return (tmp_path / "result.xlsx").read_bytes(), receipt["metadata"]

    run = asyncio.run(run_case(provider, settings(), case, _factory(), optimizer_completion_factory=completion))
    assert run.passed, run.failures
    assert len(received) == 1
    inputs = [e["schedule_yaml"] for e in run.trajectory["events"] if e["kind"] == "optimizer_input"]
    assert received[0][1] == inputs[0]
    callback = next(e["text"] for e in run.trajectory["events"] if e["kind"] == "optimizer")
    assert '"termination_reason": "optimality_proven"' in callback


@pytest.mark.parametrize("value", ["lower", "higher"])
def test_score_oracle_rejects_the_wrong_direction(value):
    case = CASE_BY_ID["optimizer-score-same-model"]
    outcome = RunOutcome(answer=json.dumps({"better": "previous", "scoreDirection": value}))
    assert grade(replace(case, tool_usage=None, optimizer_completion=""), outcome).passed == (value == "higher")


def test_existing_evidence_survives_optional_case_fields():
    validate_prompt_evidence(load_prompt_steps(), CASE_BY_ID, FIXTURE_DIGESTS)
    case = CASE_BY_ID["optimizer-score-same-model"]
    assert case_digest(case) != case_digest(replace(case, answer_json={"better": "current"}))


@pytest.mark.parametrize("answer", ['{"count": 0}', '{"count": false}', "{}", "not JSON"])
def test_structured_oracle_rejects_false_success(answer):
    case = EvalCase("count", "new-schedule", "Count", False, answer_json={"count": 0})
    assert grade(case, RunOutcome(answer=answer)).passed == (answer == '{"count": 0}')


def test_structured_oracle_uses_the_final_answer_after_streamed_commentary():
    case = EvalCase("count", "new-schedule", "Count", False, answer_json={"count": 0})
    answer = 'Initial check: {"count": 3}. Corrected the parser.\n{"count": 0}'
    assert grade(case, RunOutcome(answer=answer)).passed
    assert not grade(case, RunOutcome(answer=answer.replace('{"count": 0}', '{"count": 2}'))).passed


@pytest.mark.parametrize(
    ("actual", "passed"),
    [
        ({"Alex": {"01": "OFF", "02": "K"}}, True),
        ({"Alex": {"01": "OFF", "02": "K", "03": "OFF"}}, True),
        ({"Alex": {"01": "OFF"}}, False),
        ({"Alex": {"01": "OFF", "02": "OFF"}}, False),
    ],
)
def test_structured_oracle_checks_required_nested_fields_without_rejecting_extra_valid_assignments(actual, passed):
    case = EvalCase(
        "witness", "new-schedule", "Witness", False, answer_json={"witness": {"Alex": {"01": "OFF", "02": "K"}}}
    )
    assert grade(case, RunOutcome(answer=json.dumps({"witness": actual}))).passed == passed


def test_structured_oracle_keeps_lists_exact():
    case = EvalCase("names", "new-schedule", "Names", False, answer_json={"names": ["Alex"]})
    assert not grade(case, RunOutcome(answer='{"names": ["Alex", "Mira"]}')).passed


def test_independent_analysis_case_accepts_a_valid_extra_date_and_rejects_starting_a_job():
    case = CASE_BY_ID["optimizer-independent-request-check"]
    answer = {
        "jointlyPossible": True,
        "witness": {"Alex": {"01": "OFF", "02": "K", "03": "OFF"}, "Mira": {"02": "OFF", "03": "OFF"}},
    }
    outcome = RunOutcome(answer=json.dumps(answer))
    assert grade(case, outcome).passed
    outcome.activity = [{"kind": "tool", "name": "optimizer", "ok": True, "arguments": '{"action":"start"}'}]
    assert not grade(case, outcome).passed


@pytest.mark.parametrize(
    ("expected", "actual"),
    [
        ({"day": "OFF"}, {"day": "D"}),
        (1, 2),
        (1.5, 2.5),
        ("Ada", "Bela"),
        (True, False),
        (["01"], ["02"]),
    ],
)
def test_json_evidence_fingerprint_tracks_the_comparator_for_every_answer_shape(monkeypatch, expected, actual):
    from .ai_eval import grading

    case = EvalCase("json-answer", "new-schedule", "Return JSON.", False, answer_json={"value": expected})
    no_json = EvalCase("text-answer", "new-schedule", "Explain.", False)
    before = case_digest(case), case_digest(no_json)
    outcome = RunOutcome(answer=json.dumps({"value": actual}))
    assert not grade(case, outcome).passed

    def different_oracle(actual, expected):
        return True

    monkeypatch.setattr(grading, "_answer_json_matches", different_oracle)
    assert grade(case, outcome).passed
    assert case_digest(case) != before[0]
    assert case_digest(no_json) == before[1]


@pytest.mark.parametrize("edit_before_start", [False, True])
def test_edit_then_optimize_grades_the_submitted_snapshot(edit_before_start):
    from nurse_scheduling.loader import _load_yaml

    case = CASE_BY_ID["optimizer-edit-before-start"]
    initial = FIXTURE.read_text(encoding="utf-8")
    updated = initial.replace("weight: 11000000}", "weight: 11000000000}")
    result = grade(
        case,
        RunOutcome(
            initial=_load_yaml(initial.encode()),
            proposed=_load_yaml(updated.encode()),
            activity=[
                {"kind": "optimizer_input", "schedule_yaml": updated if edit_before_start else initial},
                {
                    "kind": "tool",
                    "name": "optimizer",
                    "ok": True,
                    "arguments": '{"action":"start","timeout_seconds":60}',
                },
            ],
        ),
    )
    assert result.passed == edit_before_start
    if not edit_before_start:
        assert any(not check.passed and check.description.startswith("optimizer input:") for check in result.checks)


@pytest.mark.parametrize("alter_source", [False, True])
def test_optimizer_only_input_must_preserve_original_policy(alter_source):
    from nurse_scheduling.loader import _load_yaml

    case = CASE_BY_ID["optimizer-start-preserves-ward"]
    source = FIXTURE.read_text(encoding="utf-8")
    submitted = source.replace("weight: 11000000000}", "weight: 11}") if alter_source else source
    result = grade(
        case,
        RunOutcome(
            initial=_load_yaml(source.encode()),
            activity=[
                {"kind": "optimizer_input", "schedule_yaml": submitted},
                {
                    "kind": "tool",
                    "name": "optimizer",
                    "ok": True,
                    "arguments": '{"action":"start","timeout_seconds":60}',
                },
            ],
        ),
    )
    assert result.passed == (not alter_source)


def test_pending_completion_uses_the_modified_input_snapshot():
    from .ai_eval.optimizer_fixtures import RESULT_SOURCES

    approved = FIXTURE.read_text(encoding="utf-8")
    pending = RESULT_SOURCES["request-audit-pending"].read_text(encoding="utf-8")
    _, approved_result = completion_result("request-audit", approved)
    _, pending_result = completion_result("request-audit-pending", pending)
    assert approved_result["source_sha256"] != pending_result["source_sha256"]
    for result, expected in [(approved_result, 1), (pending_result, 2)]:
        tier = next(row for row in result["request_audit"]["summary"] if row["weight"] == 11000000000)
        assert tier["unmet"] == expected
    with pytest.raises(ValueError, match="differs from the controlled"):
        completion_result("request-audit-pending", approved)
    with pytest.raises(ValueError, match="differs from the controlled"):
        completion_result("request-audit", pending)


def test_multi_turn_completion_prewarms_the_expected_snapshot():
    from .ai_eval.optimizer_fixtures import RESULT_SOURCES
    from .ai_eval.runner import run_all

    case = CASE_BY_ID["optimizer-result-from-pending-proposal"]
    pending = RESULT_SOURCES[case.optimizer_completion].read_text(encoding="utf-8")
    provider = ScriptedProvider(
        [ToolCallRequest((ToolCall("write", "write", json.dumps({"path": WORKSPACE_SCHEDULE, "content": pending})),))],
        [ToolCallRequest((ToolCall("start", "optimizer", '{"action":"start","timeout_seconds":60}'),))],
        [TextDelta(json.dumps(case.answer_json))],
    )
    runs = asyncio.run(run_all([case], settings(), provider, 4, _factory()))
    assert len(runs) == 1 and runs[0].passed, runs[0].failures


def test_completion_context_matches_the_submitted_proposal():
    from .ai_eval.optimizer_fixtures import RESULT_SOURCES

    case = CASE_BY_ID["optimizer-result-from-pending-proposal"]
    submitted = RESULT_SOURCES[case.optimizer_completion].read_text(encoding="utf-8")
    provider = ScriptedProvider(
        [
            ToolCallRequest(
                (ToolCall("write", "write", json.dumps({"path": WORKSPACE_SCHEDULE, "content": submitted})),)
            )
        ],
        [ToolCallRequest((ToolCall("start", "optimizer", '{"action":"start","timeout_seconds":60}'),))],
        [ToolCallRequest((ToolCall("context", "read", json.dumps({"path": WORKSPACE_RESULT_CONTEXT})),))],
        [TextDelta(json.dumps(case.answer_json))],
    )
    factory = _factory()
    run = asyncio.run(run_case(provider, settings(), case, factory))

    assert run.passed, run.failures
    result_context = json.loads(factory.created[-1].files[WORKSPACE_RESULT_CONTEXT])
    assert result_context["source_sha256"] == hashlib.sha256(submitted.encode()).hexdigest()
    promoted = [request for request in result_context["requests"] if request["preference_index"] == 4]
    assert promoted[0]["weight"] == 11_000_000_000


@pytest.mark.parametrize("review_changes_workspace", [False, True])
def test_result_review_must_not_create_another_proposal(review_changes_workspace):
    from nurse_scheduling.loader import _load_yaml

    from .ai_eval.optimizer_fixtures import RESULT_SOURCES

    case = CASE_BY_ID["optimizer-result-from-pending-proposal"]
    pending = RESULT_SOURCES[case.optimizer_completion].read_text(encoding="utf-8")
    result = grade(
        case,
        RunOutcome(
            answer=json.dumps(case.answer_json),
            initial=_load_yaml(FIXTURE.read_bytes()),
            proposed=_load_yaml(pending.encode()),
            proposal_turns=[True, review_changes_workspace],
            activity=[
                {"kind": "optimizer", "turn": 2},
                {"kind": "optimizer_input", "schedule_yaml": pending},
                {
                    "kind": "tool",
                    "name": "optimizer",
                    "ok": True,
                    "arguments": '{"action":"start","timeout_seconds":60}',
                },
            ],
        ),
    )
    assert result.passed == (not review_changes_workspace)


def test_current_request_fingerprint_ignores_unrelated_result_cli_changes(tmp_path, monkeypatch):
    from nurse_scheduling.ai import workspace
    from nurse_scheduling.ai.attachment_tools import inspect_optimizer_result as reader

    current = CASE_BY_ID["request-tier-counts-large"]
    completed = CASE_BY_ID["result-assignment-details"]
    before = case_digest(current), case_digest(completed)
    key = "/reference/tools/inspect_optimizer_result.py"
    changed = tmp_path / "inspect_optimizer_result.py"
    changed.write_text(
        workspace.REFERENCE_ATTACHMENT_TOOLS[key].read_text(encoding="utf-8") + "\n# Changed result CLI\n",
        encoding="utf-8",
    )
    monkeypatch.setitem(workspace.REFERENCE_ATTACHMENT_TOOLS, key, changed)
    assert case_digest(current) == before[0]
    assert case_digest(completed) != before[1]

    def changed_weight(value):
        return value

    monkeypatch.setattr(reader, "_weight", changed_weight)
    assert case_digest(current) != before[0]
